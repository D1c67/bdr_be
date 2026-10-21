"""Auto send: the external estimator's price, read from the estimate workbook.

Deterministic, no AI (docs/AUTO_SEND.md section 4). The estimator's workbook
carries a "Bid Recap & Summary" sheet: one row per scope group with MATERIAL
COST, LABOR COST, MATERIAL TAX, LABOR TAX, TOTAL COST, OHP (LABOR),
OHP (MATERIAL) and TOTAL PRICE columns, a TOTALS row, then a BID SUMMARY
block ending in BASE BID PRICE. Auto send prices the proposal at

    materials = material cost + material tax + material OH&P
    labor     = labor cost + labor tax + labor OH&P

from the TOTALS row, with any bid security spread over the two so materials
plus labor is exactly the BASE BID PRICE, then adds the Settings markup.

Every cell is found by its label, never by a fixed address, and the figures
are cross-checked against each other before anything is trusted: a workbook
that does not add up is refused with a plain-English reason instead of being
sent. The parse of a file is cached in `auto_send_estimates`, keyed by the
project_files row (sealed estimator files are immutable, and the sha256 is
compared when the row has one).

File choice: only sealed estimator deliverables (`estimator_deliverable` set,
`submission_round` not null) count, so internal uploads in the estimate
category (misfiled PDFs, a team member's low-voltage sheet) can never price
a proposal. Only the latest sealed round that sent estimate files counts
(owner decision: latest SEALED round only): it must hold exactly one workbook.
An older round's workbook is never used as a fallback, neither when the
newest workbook fails to parse nor when the newest round's estimate is not a
workbook (a revised estimate sent as a PDF supersedes the older price).
"""

import hashlib
import io
import logging
import re
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

from app.core.config import get_settings
from app.core.supabase_client import get_supabase
from app.services import storage

logger = logging.getLogger(__name__)

# Bumped whenever parse_recap's output or rules change, so cached rows parsed
# by older code are parsed again instead of trusted.
PARSER_VERSION = 1

RECAP_SHEET = "Bid Recap & Summary"  # plans_reader_render.RECAP_SHEET
WORKBOOK_EXTENSIONS = (".xlsx", ".xlsm")
CACHE_TABLE = "auto_send_estimates"

# Bounds on the recap sheet scan. Real recaps end well inside these (TOTALS
# between rows 7 and 14, BASE BID PRICE 14 rows below it); the caps keep a
# corrupt or inflated sheet from costing more than a few thousand cells.
HEADER_SCAN_ROWS = 30
MAX_SCAN_ROWS = 400
MAX_SCAN_COLS = 40

# Normalized header text -> component key, the eight cost columns.
COST_HEADERS = {
    "MATERIAL COST": "mc",
    "LABOR COST": "lc",
    "MATERIAL TAX": "mt",
    "LABOR TAX": "lt",
    "TOTAL COST": "tc",
    "OHP (LABOR)": "ohpl",
    "OHP (MATERIAL)": "ohpm",
    "TOTAL PRICE": "tp",
}
DESCRIPTION_HEADER = "DESCRIPTION"
TOTALS_LABEL = "TOTALS"

# BID SUMMARY labels (normalized) -> key.
SUMMARY_LABELS = {
    "TOTAL MATERIAL COST": "sum_mc",
    "MATERIAL SALES TAX": "mat_tax",
    "TOTAL LABOR COST": "sum_lc",
    "LABOR TAX": "lab_tax",
    "OVERHEAD & PROFIT (LABOR)": "ohp_labor",
    "OVERHEAD & PROFIT (MATERIAL)": "ohp_material",
    "BID SECURITY": "bid_sec",
    "BASE BID PRICE": "base",
}
# Summary rows that carry a rate in the column before the value; the rate
# only drives warnings.
RATE_ROWS = {
    "mat_tax": ("material sales tax", Decimal("0.5")),
    "lab_tax": ("labor tax", Decimal("0.5")),
    "ohp_labor": ("labor overhead and profit", Decimal("1")),
    "ohp_material": ("material overhead and profit", Decimal("1")),
}

COMPONENT_KEYS = ("mc", "lc", "mt", "lt", "tc", "ohpl", "ohpm", "tp", "bid_sec", "base")

CENT = Decimal("0.01")
CROSSCHECK_MIN_TOL = Decimal("0.05")
CROSSCHECK_REL_TOL = Decimal("0.00001")

# code -> owner-facing sentence (preflight, digest, Auto send panel). Plain
# English, no em dashes.
REASON_MESSAGES: dict[str, str] = {
    "estimate_not_received": "The estimator has not submitted an estimate yet.",
    "no_estimate_workbook": (
        "The estimator's estimate is not an Excel workbook (.xlsx), so the price cannot be read."
    ),
    "newer_estimate_not_workbook": (
        "The estimator's latest round sent an estimate that is not an Excel workbook (.xlsx), "
        "so the price cannot be read. The older round's workbook is not used. Ask the "
        "estimator to send the workbook."
    ),
    "ambiguous_workbooks": (
        "The estimator's latest round has more than one estimate workbook, so it is not clear "
        "which price to send."
    ),
    "too_large": "The estimate workbook is too large to read.",
    "download_failed": (
        "The estimate workbook could not be downloaded from storage. It will be tried again."
    ),
    "lookup_failed": (
        "The estimate could not be checked because of a server error. It will be tried again."
    ),
    "unreadable": (
        "The estimate workbook could not be opened. It may be damaged, password protected, "
        "or not a real Excel file."
    ),
    "no_recap_sheet": "The estimate workbook has no Bid Recap & Summary sheet.",
    "no_header": "The Bid Recap sheet is missing its cost column headers.",
    "no_totals_row": "The Bid Recap sheet has no TOTALS row.",
    "no_base_bid": "The Bid Recap sheet has no BASE BID PRICE.",
    "no_cached_values": (
        "The estimate workbook has no saved values. Open it in Excel, save it, and upload it again."
    ),
    "formula_error": "The Bid Recap shows an error (like #REF!) or text where a price should be.",
    "unpriced_or_zero": "The estimate is not priced (the base bid price is zero or blank).",
    "crosscheck_failed": "The Bid Recap numbers do not add up, so the price was not trusted.",
}


class EstimateError(Exception):
    """A reason the estimate cannot price an auto send. `code` is a key of
    REASON_MESSAGES; `message` is the owner-facing sentence (sometimes more
    specific than the default, for example which cross-check failed)."""

    def __init__(self, code: str, message: str | None = None):
        self.code = code
        self.message = message or REASON_MESSAGES.get(code) or code
        super().__init__(self.message)


# ---------------------------------------------------------------------------
# Money helpers
# ---------------------------------------------------------------------------


def q(value: Decimal) -> Decimal:
    """Round to cents, half up (the proposal's rounding)."""
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def _to_decimal(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise TypeError("bool is not money")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, str):
        return Decimal(value.strip())
    raise TypeError(f"not a number: {type(value).__name__}")


def apply_markup(
    materials: Decimal, labor: Decimal, pct: Decimal
) -> tuple[Decimal, Decimal, Decimal]:
    """(materials, labor, total) with `pct` percent added to each part and
    rounded to cents. The total is ALWAYS the sum of the two rounded parts, so
    the three numbers on the proposal add up."""
    factor = Decimal(1) + _to_decimal(pct) / Decimal(100)
    m = q(_to_decimal(materials) * factor)
    lab = q(_to_decimal(labor) * factor)
    return m, lab, m + lab


def _money(value: Decimal | None) -> str | None:
    return None if value is None else str(q(value))


# ---------------------------------------------------------------------------
# Workbook reading
# ---------------------------------------------------------------------------


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip().upper()


def _norm_label(value: Any) -> str:
    """Summary label text without the trailing '@' or ':' some rows carry."""
    return _norm(value).rstrip("@:").strip()


def _open_workbook(xlsx_bytes: bytes, *, data_only: bool):
    import openpyxl

    try:
        return openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=data_only, read_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, KeyError, ...
        raise EstimateError("unreadable") from exc


def _recap_title(sheetnames: list[str]) -> str | None:
    if RECAP_SHEET in sheetnames:
        return RECAP_SHEET
    return next((n for n in sheetnames if "recap" in (n or "").lower()), None)


def _sheet_grid(wb, title: str) -> dict[tuple[int, int], Any]:
    """{(row, col): value} of the non-empty cells in the bounded scan area."""
    ws = wb[title]
    if not hasattr(ws, "iter_rows"):  # a chartsheet named like a recap
        raise EstimateError("no_recap_sheet")
    grid: dict[tuple[int, int], Any] = {}
    for r, row in enumerate(
        ws.iter_rows(max_row=MAX_SCAN_ROWS, max_col=MAX_SCAN_COLS, values_only=True), start=1
    ):
        for c, v in enumerate(row, start=1):
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            grid[(r, c)] = v
    return grid


class _Recap:
    """The recap sheet's values, with the formula view loaded only when a
    cell we need is empty (to tell 'never calculated' from 'blank')."""

    def __init__(self, xlsx_bytes: bytes, title: str, values: dict[tuple[int, int], Any]):
        self._bytes = xlsx_bytes
        self._title = title
        self.values = values
        self._formulas: dict[tuple[int, int], Any] | None = None

    def has_formula(self, cell: tuple[int, int]) -> bool:
        if self._formulas is None:
            wb = _open_workbook(self._bytes, data_only=False)
            try:
                self._formulas = _sheet_grid(wb, self._title)
            except EstimateError:
                raise
            except Exception as exc:
                raise EstimateError("unreadable") from exc
            finally:
                wb.close()
        v = self._formulas.get(cell)
        return isinstance(v, str) and v.startswith("=")

    def number(
        self, cell: tuple[int, int], *, blank: Decimal | None = Decimal(0)
    ) -> Decimal | None:
        """The cell as a Decimal. Error values and text raise formula_error; an
        empty cell holding a formula raises no_cached_values (the workbook was
        written by a tool that does not calculate); a truly blank cell is
        `blank`."""
        v = self.values.get(cell)
        if v is None:
            if self.has_formula(cell):
                raise EstimateError("no_cached_values")
            return blank
        if isinstance(v, str) and v.strip().startswith("#"):
            raise EstimateError("formula_error")
        if isinstance(v, str):
            v = v.replace(",", "").replace("$", "")
        try:
            d = _to_decimal(v)
        except (TypeError, InvalidOperation):
            raise EstimateError("formula_error") from None
        if not d.is_finite():
            raise EstimateError("formula_error")
        return d


def _find_header(values: dict[tuple[int, int], Any]) -> tuple[int, dict[str, int], int]:
    """(header row, {key: column}, description column)."""
    by_row: dict[int, dict[str, int]] = {}
    desc: dict[int, int] = {}
    for (r, c), v in values.items():
        if r > HEADER_SCAN_ROWS or not isinstance(v, str):
            continue
        t = _norm(v)
        if t in COST_HEADERS:
            by_row.setdefault(r, {}).setdefault(COST_HEADERS[t], c)
        elif t == DESCRIPTION_HEADER:
            desc.setdefault(r, c)
    for r in sorted(by_row):
        cols = by_row[r]
        if len(cols) == len(COST_HEADERS):
            # No DESCRIPTION header: the label column is the one left of the
            # first cost column (where every real and generated recap has it).
            dcol = desc.get(r) or (min(cols.values()) - 1)
            if dcol < 1:
                break
            return r, cols, dcol
    raise EstimateError("no_header")


def _label_rows(
    values: dict[tuple[int, int], Any], dcol: int, start_row: int
) -> tuple[int | None, dict[str, int]]:
    """(TOTALS row, {summary key: row}) from the description column, below
    the header. Summary labels are only looked for below TOTALS."""
    totals: int | None = None
    summary: dict[str, int] = {}
    rows = sorted(r for (r, c) in values if c == dcol and r > start_row)
    for r in rows:
        v = values[(r, dcol)]
        if not isinstance(v, str):
            continue
        if totals is None:
            if _norm(v) == TOTALS_LABEL:
                totals = r
            continue
        key = SUMMARY_LABELS.get(_norm_label(v))
        if key:
            summary.setdefault(key, r)
    return totals, summary


def _value_col(recap: _Recap, base_row: int, dcol: int) -> int | None:
    """The summary value column: the first cell right of the BASE BID PRICE
    label with a value (or, in a workbook without cached values, a formula).
    Rate rows (taxes, OH&P, bid security) keep their rate one column left
    of it. None when the base bid cell is blank."""
    for c in range(dcol + 1, MAX_SCAN_COLS + 1):
        if (base_row, c) in recap.values:
            return c
    for c in range(dcol + 1, MAX_SCAN_COLS + 1):
        if recap.has_formula((base_row, c)):
            return c
    return None


def _crosscheck(name: str, a: Decimal, b: Decimal, tol: Decimal) -> None:
    if abs(a - b) > tol:
        raise EstimateError(
            "crosscheck_failed",
            f"The Bid Recap numbers do not add up ({name}: {q(a):,} vs {q(b):,}), "
            "so the price was not trusted.",
        )


def parse_recap(xlsx_bytes: bytes) -> dict[str, Any]:
    """Read the estimator price from a workbook's Bid Recap. Pure (no I/O
    beyond the bytes); raises EstimateError with a REASON_MESSAGES code.

    Returns {"sheet", "components": {mc, lc, mt, lt, tc, ohpl, ohpm, tp,
    bid_sec, base} (Decimal, as read), "materials", "labor", "base" (Decimal,
    cents, materials + labor == base exactly), "warnings": [str]}."""
    wb = _open_workbook(xlsx_bytes, data_only=True)
    try:
        title = _recap_title(list(wb.sheetnames))
        if title is None:
            raise EstimateError("no_recap_sheet")
        try:
            values = _sheet_grid(wb, title)
        except EstimateError:
            raise
        except Exception as exc:
            raise EstimateError("unreadable") from exc
    finally:
        wb.close()
    recap = _Recap(xlsx_bytes, title, values)

    header_row, cols, dcol = _find_header(values)
    totals_row, summary = _label_rows(values, dcol, header_row)
    if totals_row is None:
        raise EstimateError("no_totals_row")
    if "base" not in summary:
        raise EstimateError("no_base_bid")

    comp: dict[str, Decimal] = {}
    # Every cost cell is read before any check, so an error or a missing
    # cached value anywhere on the TOTALS row wins over a misleading
    # "does not add up".
    for key in COST_HEADERS.values():
        comp[key] = recap.number((totals_row, cols[key]))

    base_row = summary["base"]
    vcol = _value_col(recap, base_row, dcol)
    if vcol is None:
        raise EstimateError("unpriced_or_zero")
    base = recap.number((base_row, vcol))
    bid_sec = Decimal(0)
    if "bid_sec" in summary:
        bid_sec = recap.number((summary["bid_sec"], vcol))
    sum_mc = recap.number((summary["sum_mc"], vcol), blank=None) if "sum_mc" in summary else None
    sum_lc = recap.number((summary["sum_lc"], vcol), blank=None) if "sum_lc" in summary else None
    comp["bid_sec"] = bid_sec
    comp["base"] = base

    warnings: list[str] = []
    for key, (what, ceiling) in RATE_ROWS.items():
        row = summary.get(key)
        if row is None or vcol - 1 <= dcol:
            continue
        try:
            rate = recap.number((row, vcol - 1), blank=None)
        except EstimateError:
            rate = None
        if rate is not None and not (Decimal(0) <= rate <= ceiling):
            warnings.append(f"The {what} rate on the Bid Recap is unusual ({rate}).")

    if base <= 0 or comp["tp"] <= 0:
        raise EstimateError("unpriced_or_zero")

    materials_raw = comp["mc"] + comp["mt"] + comp["ohpm"]
    labor_raw = comp["lc"] + comp["lt"] + comp["ohpl"]

    tol = max(CROSSCHECK_MIN_TOL, base * CROSSCHECK_REL_TOL)
    _crosscheck(
        "costs vs TOTAL COST", comp["mc"] + comp["lc"] + comp["mt"] + comp["lt"], comp["tc"], tol
    )
    _crosscheck(
        "TOTAL COST plus overhead and profit vs TOTAL PRICE",
        comp["tc"] + comp["ohpl"] + comp["ohpm"],
        comp["tp"],
        tol,
    )
    _crosscheck("materials plus labor vs TOTAL PRICE", materials_raw + labor_raw, comp["tp"], tol)
    _crosscheck("TOTAL PRICE plus BID SECURITY vs BASE BID PRICE", comp["tp"] + bid_sec, base, tol)
    if sum_mc is not None:
        _crosscheck("TOTAL MATERIAL COST vs the TOTALS row", sum_mc, comp["mc"], tol)
    if sum_lc is not None:
        _crosscheck("TOTAL LABOR COST vs the TOTALS row", sum_lc, comp["lc"], tol)
    if materials_raw < 0 or labor_raw < 0:
        raise EstimateError(
            "crosscheck_failed",
            "The Bid Recap has a negative materials or labor total, so the price was not trusted.",
        )

    # Bid security spread in proportion to materials and labor so the two
    # always sum to the base bid; labor takes the rounding remainder.
    base_out = q(base)
    mat_share = bid_sec * materials_raw / comp["tp"]
    materials = q(materials_raw + mat_share)
    labor = base_out - materials
    if labor < 0 or materials < 0:
        raise EstimateError(
            "crosscheck_failed",
            "The Bid Recap has a negative materials or labor total, so the price was not trusted.",
        )
    if bid_sec != 0:
        warnings.append(
            f"Bid security of {q(bid_sec):,} is included, split between materials and labor."
        )
    if materials == 0:
        warnings.append("The estimate has no material cost (labor only).")
    if labor == 0:
        warnings.append("The estimate has no labor cost (materials only).")

    return {
        "sheet": title,
        "components": comp,
        "materials": materials,
        "labor": labor,
        "base": base_out,
        "warnings": warnings,
    }


# ---------------------------------------------------------------------------
# Picking the file
# ---------------------------------------------------------------------------

_FILE_COLUMNS = (
    "id, filename, storage_path, size_bytes, submission_round, uploaded_by, sha256, created_at, "
    "impersonated_by"
)


def _is_workbook(filename: Any) -> bool:
    return str(filename or "").strip().lower().endswith(WORKBOOK_EXTENSIONS)


def pick_estimate_file(sb, project_id: str) -> dict[str, Any]:
    """The sealed estimator workbook that prices the project: the one
    workbook of the latest sealed round that sent estimate files. Raises
    EstimateError: estimate_not_received (no sealed estimate at all),
    no_estimate_workbook (sealed estimate files, none a workbook),
    newer_estimate_not_workbook (the latest round's estimate files hold no
    workbook, so an older round's price would be stale), ambiguous_workbooks
    (more than one workbook in that round). The returned row's
    "newer_round_without_workbook" is always None (kept for the wire shape)."""
    rows = (
        sb.table("project_files")
        .select(_FILE_COLUMNS)
        .eq("project_id", project_id)
        .eq("category", "estimate")
        .eq("estimator_deliverable", True)
        .not_.is_("submission_round", "null")
        .execute()
        .data
        or []
    )
    rows = [r for r in rows if r.get("submission_round") is not None]
    if not rows:
        raise EstimateError("estimate_not_received")
    books = [r for r in rows if _is_workbook(r.get("filename"))]
    if not books:
        raise EstimateError("no_estimate_workbook")
    top = max(int(r["submission_round"]) for r in books)
    newest_any = max(int(r["submission_round"]) for r in rows)
    if newest_any > top:
        raise EstimateError(
            "newer_estimate_not_workbook",
            f"The estimator's latest round (round {newest_any}) sent an estimate that is not "
            f"an Excel workbook (.xlsx), so the price cannot be read. The round {top} workbook "
            f"is not used. Ask the estimator to send the workbook for round {newest_any}.",
        )
    in_top = [r for r in books if int(r["submission_round"]) == top]
    if len(in_top) > 1:
        raise EstimateError("ambiguous_workbooks")
    out = dict(in_top[0])
    out["newer_round_without_workbook"] = None
    return out


# ---------------------------------------------------------------------------
# load_estimate: pick + cache + download + parse, never raises
# ---------------------------------------------------------------------------


def _parse_payload(parsed: dict[str, Any]) -> dict[str, Any]:
    """JSON-safe copy of a parse_recap result (money as cent strings)."""
    return {
        "parser_version": PARSER_VERSION,
        "sheet": parsed["sheet"],
        "components": {k: _money(parsed["components"].get(k)) for k in COMPONENT_KEYS},
        "materials": _money(parsed["materials"]),
        "labor": _money(parsed["labor"]),
        "base": _money(parsed["base"]),
        "warnings": list(parsed["warnings"]),
    }


def _result(
    *,
    ok: bool,
    code: str | None = None,
    message: str | None = None,
    file: dict[str, Any] | None = None,
    estimator_name: str | None = None,
    parsed: dict[str, Any] | None = None,
) -> dict[str, Any]:
    f = file or {}
    p = parsed or {}
    warnings = list(p.get("warnings") or [])
    newer = f.get("newer_round_without_workbook")
    if newer is not None and f.get("submission_round") is not None:
        warnings.insert(
            0,
            f"Round {newer} has no estimate workbook, so the round {f['submission_round']} "
            "workbook is used.",
        )
    return {
        "ok": ok,
        "code": code,
        "message": message,
        "file_id": f.get("id"),
        "filename": f.get("filename"),
        "round": f.get("submission_round"),
        "uploaded_by": f.get("uploaded_by"),
        # Who really priced it under dev "View as" (0159), and when the
        # workbook was uploaded: auto_send's dev-estimator check reads both.
        "impersonated_by": f.get("impersonated_by"),
        "uploaded_at": f.get("created_at"),
        "estimator_name": estimator_name,
        "sheet": p.get("sheet"),
        "components": dict(p.get("components") or {}),
        "materials": p.get("materials"),
        "labor": p.get("labor"),
        "base": p.get("base"),
        "warnings": warnings,
        "newer_round_without_workbook": newer,
    }


def _failed(code: str, message: str | None = None, **kw: Any) -> dict[str, Any]:
    return _result(ok=False, code=code, message=message or REASON_MESSAGES.get(code), **kw)


def _estimator_name(sb, user_id: str | None) -> str | None:
    if not user_id:
        return None
    try:
        rows = (
            sb.table("profiles").select("full_name").eq("id", user_id).limit(1).execute().data or []
        )
    except Exception:
        logger.warning("auto send: estimator name lookup failed for %s", user_id, exc_info=True)
        return None
    return (rows[0].get("full_name") or None) if rows else None


def _cached(sb, file: dict[str, Any]) -> dict[str, Any] | None:
    try:
        rows = (
            sb.table(CACHE_TABLE)
            .select("sha256, ok, code, message, result")
            .eq("estimate_file_id", file["id"])
            .limit(1)
            .execute()
            .data
            or []
        )
    except Exception:
        logger.warning("auto send: estimate cache read failed for %s", file["id"], exc_info=True)
        return None
    if not rows:
        return None
    row = rows[0]
    if (row.get("result") or {}).get("parser_version") != PARSER_VERSION:
        return None
    if file.get("sha256") and row.get("sha256") != file["sha256"]:
        return None
    return row


def _store(sb, project_id: str, file: dict[str, Any], sha256: str | None, row: dict) -> None:
    try:
        sb.table(CACHE_TABLE).upsert(
            {
                "estimate_file_id": file["id"],
                "project_id": project_id,
                "sha256": sha256,
                "parsed_at": datetime.now(timezone.utc).isoformat(),
                **row,
            },
            on_conflict="estimate_file_id",
        ).execute()
    except Exception:
        logger.warning("auto send: estimate cache write failed for %s", file["id"], exc_info=True)


def load_estimate(project_id: str, *, sb=None) -> dict[str, Any]:
    """The project's EstimateResult (docs/AUTO_SEND.md section 4). Never
    raises: every failure is {"ok": False, "code", "message", ...}.

    Transient failures (database, storage download) are not cached; parse
    outcomes, good or bad, are cached per file."""
    try:
        sb = sb or get_supabase()
        try:
            file = pick_estimate_file(sb, project_id)
        except EstimateError as e:
            return _failed(e.code, e.message)
    except Exception:
        logger.exception("auto send: estimate pick failed for project %s", project_id)
        return _failed("lookup_failed")

    name = _estimator_name(sb, file.get("uploaded_by"))
    ctx = {"file": file, "estimator_name": name}
    try:
        hit = _cached(sb, file)
        if hit is not None:
            if hit.get("ok"):
                return _result(ok=True, parsed=hit.get("result") or {}, **ctx)
            return _failed(hit.get("code") or "unreadable", hit.get("message"), **ctx)

        size = file.get("size_bytes")
        if size is not None and int(size) > get_settings().boq_max_bytes:
            return _failed("too_large", **ctx)
        try:
            data = storage.download_file(file["storage_path"])
        except Exception:
            logger.warning(
                "auto send: estimate download failed for %s", file.get("id"), exc_info=True
            )
            return _failed("download_failed", **ctx)
        if len(data) > get_settings().boq_max_bytes:
            return _failed("too_large", **ctx)
        sha = file.get("sha256") or hashlib.sha256(data).hexdigest()

        try:
            payload = _parse_payload(parse_recap(data))
        except EstimateError as e:
            _store(
                sb,
                project_id,
                file,
                sha,
                {
                    "ok": False,
                    "code": e.code,
                    "message": e.message,
                    "result": {"parser_version": PARSER_VERSION},
                },
            )
            return _failed(e.code, e.message, **ctx)
        _store(
            sb,
            project_id,
            file,
            sha,
            {"ok": True, "code": None, "message": None, "result": payload},
        )
        return _result(ok=True, parsed=payload, **ctx)
    except Exception:
        logger.exception("auto send: estimate load failed for project %s", project_id)
        return _failed("lookup_failed", **ctx)
