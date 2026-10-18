"""Deliverable files for AI Estimating (docs/PLANS_READER.md 3.6 and 6).

The main agent returns JSON specs (services/plans_reader_schemas.py); this
module turns them into the files the estimators already send, so a generated
BOQ or estimate opens looking exactly like one an estimator built by hand.
The layout constants below were measured on the 2026 sent files in the
2026-10-01 production backup (Edikted, Traffic Signal Improvements, Five Star
Tax Resolution, YMCA Torrey Pines), not copied from memory:

- `BOQ - COUNTS.xlsx`: a blank 'Sheet1', then 'SHEET' (one area) or one sheet
  per area. B2 holds 'NAME\\nADDRESS'; row 4 is the header (fill FF2DC8FF);
  column A carries the =IF(C<>"",1+MAX(...),"") serial on every table row
  except fitting rows, which sit unnumbered and right-aligned under their
  conduit. Quantities include the 10% wastage on FT conductor and conduit
  rows, and the fitting rows are generated from the post-wastage length with
  the same ratios the estimators use.
- `<PROJECT> - ELECTRICAL ESTIMATE.xlsx`: 'Sheet1', 'Detailed Estimate' (or
  one sheet per area) and 'Bid Recap & Summary'. Quantities are pre-wastage
  with the wastage in column E and F = D*(1+E). The estimate never lists
  fitting rows (none of the sent estimates do). Every pricing column header
  is present and every pricing cell is EMPTY: AI Estimating never prices.
- `Trench Markup.pdf` / `Sawcut Markup.pdf`: the referenced source sheets
  with red /Line annotations (width 7, as in the sent markups) and one
  /FreeText label per sheet, each with an appearance stream so every viewer
  draws them.
- `Marked Plans.pdf`: letter landscape pages with the sheet image, colored
  count dots or polylines and a legend box, built with Pillow.
- `Deep Brief.docx` (python-docx) and its PDF through the configured office
  converter when there is one.

Every string written to a file is cleaned: em and en dashes become hyphens
(house style) and text cells that would start a spreadsheet formula are
stored as plain text with Excel's quote prefix (the same CWE-1236 guard as
cpr_generation._formula_safe, without a visible apostrophe), because the
text comes from a model reading third-party documents. Raised errors are
app-authored ValueErrors so llm_errors and the routers can show them as
they are. Every function here is synchronous and CPU or network bound
(brief_pdf_bytes calls the office converter): call them from a worker or
through asyncio.to_thread, never directly inside `async def`.
"""

from __future__ import annotations

import datetime
import io
import logging
import math
import re
from collections.abc import Mapping
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.worksheet.worksheet import Worksheet

from app.services.plans_reader_schemas import (
    ESTIMATE_ALLOWANCES,
    LOW_VOLTAGE_NOTE,
    WIRING_WASTAGE_NOTE,
)

log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Shared text and number helpers
# ---------------------------------------------------------------------------

WASTAGE_FACTOR = 1.1
WASTAGE_RATE = 0.1

# U+2010..U+2015 (hyphen, figure dash, en dash, em dash, bar) and U+2212 minus.
_DASHES = re.compile(r"[\u2010-\u2015\u2212]")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _clean(text: Any) -> str:
    """House-style text for a file: no em or en dashes, no control characters
    (openpyxl refuses them), stripped."""
    if text is None:
        return ""
    s = _DASHES.sub("-", str(text))
    s = _CTRL.sub("", s)
    return s.strip()


# Lead characters that make Excel or Sheets treat typed text as a formula.
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")


def _set_text(cell, text: Any):
    """Write cleaned text into an openpyxl cell (None when empty). Text that
    would start a formula ('=HYPERLINK(...)', but also '-48V DC' or '+/- 5%')
    is forced to a string cell with Excel's quote prefix: it can never be
    evaluated, even after a user edits the cell, and it displays without the
    literal apostrophe cpr_generation._formula_safe would add."""
    s = _clean(text)
    cell.value = s or None
    if s.startswith(_FORMULA_LEADS):
        cell.data_type = "s"
        cell.quotePrefix = True
    return cell


def _text_cell(ws: Worksheet, row: int, column: int, text: Any):
    """`ws.cell(row, column)` holding `text` written through _set_text."""
    return _set_text(ws.cell(row=row, column=column), text)


def _num(value: Any, what: str) -> float:
    """A finite float from a spec quantity, else an app-authored ValueError
    naming `what`."""
    try:
        n = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{what} needs a numeric quantity.") from None
    if math.isnan(n) or math.isinf(n):
        raise ValueError(f"{what} needs a finite quantity.")
    return n


def _ceil(x: float) -> int:
    """ceil that ignores float noise (110.00000000001 / 10 is 11, not 12)."""
    return int(math.ceil(round(x, 6)))


def _round_half_up(x: float) -> int:
    """Round .5 up (Python's round() is banker's rounding), ignoring float
    noise, the way Excel's ROUND treats the elbow counts."""
    return int(math.floor(round(x, 6) + 0.5))


def boq_quantity(item: dict[str, Any]) -> float:
    """The quantity the BOQ shows: the spec quantity plus 10% when the row
    carries wastage and is measured in feet (conductors, cables, conduits),
    rounded to 2 decimals. Every other row shows the spec quantity."""
    q = _num(item.get("quantity"), f"BOQ item '{_clean(item.get('description'))[:60]}'")
    if item.get("wastage_applies") and _clean(item.get("unit")).upper() == "FT":
        return round(q * WASTAGE_FACTOR, 2)
    return round(q, 2)


def estimate_wastage(item: dict[str, Any]) -> float:
    """Column E of the estimate: 10% on FT wastage rows, 0 on everything else
    (matches the sent estimates, where saw cut and ground wire carry 0%)."""
    if item.get("wastage_applies") and _clean(item.get("unit")).upper() == "FT":
        return WASTAGE_RATE
    return 0.0


# ---------------------------------------------------------------------------
# Conduit fittings (docs 3.6 "fitting rows with the rules observed")
# ---------------------------------------------------------------------------

# '3/4"C EMT', '1-1/4"C EMT', '2"C PVC SCH 40', '4"C PVC WITH PULL STRING ...'
_CONDUIT_RE = re.compile(
    r'^\s*(?P<size>[0-9][0-9/\- ]*?)\s*"\s*C\s+(?P<type>EMT|RGS|PVC)\b', re.IGNORECASE
)

ANCHOR_DESC = "#8 To #10 X 7/8 Plas Anchor (3/16)"
SCREW_DESC = "#10 X 1 P/H Self-Tap Screw"
SQ_BOX_DESC = "SQ Box COMB KO With Blank Cover"
OTHER_CONDUIT_ACCESSORIES = "OTHER CONDUIT ACCESSORIES"


def parse_conduit(description: str) -> tuple[str, str] | None:
    """(size, type) for a conduit row such as '3/4"C EMT', else None. Only
    rows whose description contains '"C ' followed by EMT, RGS or PVC count,
    the way the estimators write conduit rows."""
    m = _CONDUIT_RE.match(_clean(description))
    if not m:
        return None
    return m.group("size").strip(), m.group("type").upper()


def conduit_fittings(description: str, length_ft: float) -> tuple[list[tuple[str, int, str]], int]:
    """Fitting rows for one conduit run of `length_ft` (the post-wastage BOQ
    length) and the 1-H strap count it contributes to OTHER CONDUIT
    ACCESSORIES. Ratios measured on the sent BOQs:

    - EMT / RGS: compression connector ceil(L/50), coupling ceil(L/10),
      1-H strap = couplings.
    - PVC (underground): elbow round(0.75 * adapters), locknut = adapters,
      coupling ceil(L/10), female adapter ceil(L/50), in that order.
    """
    parsed = parse_conduit(description)
    if parsed is None or length_ft <= 0:
        return [], 0
    size, kind = parsed
    if kind in ("EMT", "RGS"):
        conn = _ceil(length_ft / 50)
        coupling = _ceil(length_ft / 10)
        rows = [
            (f'{size}"C {kind} Compression Conn Stl', conn, "EA"),
            (f'{size}"C {kind} Compression Coupling Stl', coupling, "EA"),
            (f'{size}"C {kind} 1-H Strap - Steel', coupling, "EA"),
        ]
        return rows, coupling
    adapters = _ceil(length_ft / 50)
    rows = [
        (f'{size}" ELBOW 90 DEG - RMC - GALV', max(1, _round_half_up(0.75 * adapters)), "EA"),
        (f'{size}" LOCKNUT - STEEL', adapters, "EA"),
        (f'{size}" COUPLING - PVC', _ceil(length_ft / 10), "EA"),
        (f'{size}" ADAPTER FEM - PVC', adapters, "EA"),
    ]
    return rows, 0


def other_conduit_accessories(straps: int) -> list[tuple[str, int, str]]:
    """Anchors and screws ceil(straps/2), SQ boxes ceil(anchors/5)."""
    if straps <= 0:
        return []
    anchors = _ceil(straps / 2)
    return [
        (ANCHOR_DESC, anchors, "EA"),
        (SCREW_DESC, anchors, "EA"),
        (SQ_BOX_DESC, _ceil(anchors / 5), "EA"),
    ]


# ---------------------------------------------------------------------------
# Spec walking shared by the BOQ and the estimate
# ---------------------------------------------------------------------------


def _group_heading(name: str, note: str | None) -> str:
    """'WIRING' plus the 10% wastage note, 'LOW VOLTAGE WORK' plus the
    outlets note, or the group name with its own note on a second line."""
    n = _clean(name)
    note_s = _clean(note)
    if not note_s:
        if n.upper() == "WIRING":
            note_s = WIRING_WASTAGE_NOTE
        elif n.upper() == "LOW VOLTAGE WORK":
            note_s = LOW_VOLTAGE_NOTE
    if note_s and note_s.upper() not in n.upper():
        return f"{n}\n{note_s}"
    return n


def _subheading(name: str, note: str | None) -> str:
    """A subgroup or section heading with its note on a second line, unless
    the name already carries the note."""
    n = _clean(name)
    note_s = _clean(note)
    if note_s and note_s.upper() not in n.upper():
        return f"{n}\n{note_s}"
    return n


def _item_description(item: dict[str, Any]) -> str:
    """Description with the inline caveat ('(ASSUMED)', ...) and the material
    note on its own line in parentheses, the way the sent BOQs carry them."""
    desc = _clean(item.get("description"))
    caveat = _clean(item.get("caveat"))
    if caveat and caveat.upper() not in desc.upper():
        desc = f"{desc} {caveat}"
    note = _clean(item.get("material_note"))
    if note:
        if not note.startswith("("):
            note = f"({note})"
        desc = f"{desc}\n{note}"
    return desc


def _items(subgroup: dict[str, Any]) -> list[dict[str, Any]]:
    """The subgroup's item dicts that have a description (the model
    sometimes emits empty placeholder rows)."""
    return [
        i
        for i in (subgroup.get("items") or [])
        if isinstance(i, dict) and _clean(i.get("description"))
    ]


def _estimate_items(subgroup: dict[str, Any]) -> list[dict[str, Any]]:
    """The rows the estimate lists for a subgroup: no OTHER CONDUIT
    ACCESSORIES subgroup and no is_accessory rows (fittings live in the BOQ
    only, as in every sent estimate)."""
    if _clean(subgroup.get("name")).upper() == OTHER_CONDUIT_ACCESSORIES:
        return []
    return [i for i in _items(subgroup) if not i.get("is_accessory")]


def _group_has_items(group: dict[str, Any], *, estimate: bool = False) -> bool:
    """True when any subgroup has rows to write (estimate rows when
    `estimate`), so empty groups get no heading, subtotal or recap line."""
    pick = _estimate_items if estimate else _items
    return any(pick(sg) for sg in group.get("subgroups") or [] if isinstance(sg, dict))


def _section_has_items(section: dict[str, Any], *, estimate: bool = False) -> bool:
    """True when any group of the section has rows to write."""
    return any(
        _group_has_items(g, estimate=estimate)
        for g in section.get("groups") or []
        if isinstance(g, dict)
    )


def _areas(boq_spec: dict[str, Any]) -> list[dict[str, Any]]:
    """The spec's area dicts, or an app-authored ValueError when there are
    none."""
    if not isinstance(boq_spec, dict):
        raise ValueError("The BOQ spec must be an object.")
    areas = [a for a in boq_spec.get("areas") or [] if isinstance(a, dict)]
    if not areas:
        raise ValueError("The BOQ spec has no areas to render.")
    return areas


_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
_RESERVED_SHEETS = {"sheet1", "bid recap & summary", "chart1"}


def sheet_title(name: str, used: set[str]) -> str:
    """An Excel-safe, unique sheet name: no []:*?/\\, no leading or trailing
    apostrophe, at most 31 characters, never one of the fixed sheet names."""
    base = _SHEET_BAD.sub(" ", _clean(name))
    base = re.sub(r"\s+", " ", base).strip().strip("'").strip() or "AREA"
    base = base[:31].rstrip()
    title = base
    n = 2
    while title.lower() in used or title.lower() in _RESERVED_SHEETS:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)].rstrip() + suffix
        n += 1
    used.add(title.lower())
    return title


def _area_titles(areas: list[dict[str, Any]], single_name: str) -> list[str]:
    """Sheet names: the fixed single-area name, or one unique sanitized name
    per area."""
    if len(areas) == 1:
        return [single_name]
    used: set[str] = set()
    return [sheet_title(a.get("area") or f"AREA {i + 1}", used) for i, a in enumerate(areas)]


# ---------------------------------------------------------------------------
# openpyxl styling (values measured on the sent files)
# ---------------------------------------------------------------------------

HEADER_FILL = "FF2DC8FF"
LEVEL0_FILL = "FFFFFF00"
SUBTOTAL_FILL = "FF6DD9FF"
RECAP_TITLE_FILL = "FF00B0F0"
NOTE_FILL = "FFFFFF00"
PROJECT_RED = "FFC00000"

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _solid(rgb: str) -> PatternFill:
    """A solid fill in an ARGB color."""
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


def _theme(theme: int, tint: float) -> PatternFill:
    """A solid fill in a theme color with a tint, as the sent files use."""
    return PatternFill(fill_type="solid", fgColor=Color(theme=theme, tint=tint))


GROUP_FILL = (3, 0.3999755851924192)
LOW_VOLTAGE_FILL = (6, 0.5999938962981048)
SUBGROUP_FILL = (3, 0.7999816888943144)
ACCESSORY_HEADING_FILL = (4, 0.5999938962981048)
DATE_FILL = (6, 0.5999938962981048)


def _group_fill(name: str) -> PatternFill:
    if _clean(name).upper().startswith("LOW VOLTAGE WORK"):
        return _theme(*LOW_VOLTAGE_FILL)
    return _theme(*GROUP_FILL)


def _style(
    cell,
    *,
    bold: bool = False,
    size: float | None = None,
    h: str | None = None,
    v: str | None = None,
    wrap: bool = False,
    border: bool = True,
    fill: PatternFill | None = None,
    color: str | None = None,
    fmt: str | None = None,
) -> None:
    cell.font = Font(name="Calibri", bold=bold, size=size or 11, color=color)
    cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if border:
        cell.border = _BORDER
    if fill is not None:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt


def _qty_format(value: float) -> str:
    """'0' for whole quantities (the sent files' format) and '0.##' for
    fractional ones, so 0.4 CY never displays as 0."""
    return "0" if float(value).is_integer() else "0.##"


def _line_count(text: str, width_chars: float) -> int:
    """Approximate wrapped line count, used for row heights."""
    lines = 0
    for part in text.split("\n"):
        lines += max(1, math.ceil(len(part) / max(1.0, width_chars)))
    return lines


def _workbook_bytes(wb: Workbook) -> bytes:
    """The workbook serialized to xlsx bytes."""
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# BOQ - COUNTS.xlsx
# ---------------------------------------------------------------------------

BOQ_WIDTHS = {"A": 6.453125, "B": 68.1796875, "C": 9.453125, "D": 6.81640625, "E": 8.6328125}
BOQ_HEADER_ROW = 4
BOQ_FIRST_ROW = 5


def _boq_title(boq_spec: dict[str, Any]) -> str:
    title = _clean(boq_spec.get("title"))
    sub = _clean(boq_spec.get("subtitle"))
    return f"{title}\n{sub}" if sub else title


class _BoqSheet:
    """Writes one BOQ sheet row by row, tracking the serial formula."""

    def __init__(self, ws: Worksheet) -> None:
        self.ws = ws
        self.row = BOQ_FIRST_ROW

    def _serial(self, r: int) -> None:
        start = "A$4:A4" if r == BOQ_FIRST_ROW else f"A${BOQ_FIRST_ROW}:A{r - 1}"
        c = self.ws.cell(row=r, column=1, value=f'=IF(C{r}<>"",1+MAX({start}),"")')
        _style(c, bold=True, size=12, h="center", fmt="0")

    def blank(self) -> None:
        self._serial(self.row)
        self.row += 1

    def heading(self, text: str, *, fill: PatternFill, center: bool, numbered: bool = True) -> None:
        r = self.row
        if numbered:
            self._serial(r)
        c = _text_cell(self.ws, r, 2, text)
        _style(c, bold=True, h="center" if center else "left", wrap=True, fill=fill)
        if "\n" in text:
            self.ws.row_dimensions[r].height = 15 * _line_count(text, 62)
        self.row += 1

    def item(self, description: str, quantity: float, unit: str) -> None:
        r = self.row
        self._serial(r)
        self._row_cells(r, description, quantity, unit, accessory=False)
        self.row += 1

    def accessory(self, description: str, quantity: float, unit: str) -> None:
        r = self.row
        self._row_cells(r, description, quantity, unit, accessory=True)
        self.row += 1

    def _row_cells(
        self, r: int, description: str, quantity: float, unit: str, *, accessory: bool
    ) -> None:
        b = _text_cell(self.ws, r, 2, description)
        _style(b, h="right" if accessory else None, wrap=True)
        q = quantity if not float(quantity).is_integer() else int(quantity)
        c = self.ws.cell(row=r, column=3, value=q)
        _style(c, h="center", wrap=True, fmt=_qty_format(q))
        d = _text_cell(self.ws, r, 4, _clean(unit).upper())
        _style(d, h="center")
        lines = _line_count(description, 62)
        if lines > 1:
            self.ws.row_dimensions[r].height = 15 * lines

    def note(self, text: str) -> None:
        r = self.row
        body = _clean(text)
        value = body if body.upper().startswith("NOTE") else f"NOTE:\n{body}"
        c = _text_cell(self.ws, r, 2, value)
        _style(c, wrap=True)
        self.ws.row_dimensions[r].height = 15 * _line_count(value, 62)
        self.row += 1


def _write_boq_subgroup(sheet: _BoqSheet, group: dict[str, Any], sg_index: int) -> None:
    """One subgroup: optional heading, items with their fitting rows, the
    OTHER CONDUIT ACCESSORIES block, then a blank row."""
    subgroups = group.get("subgroups") or []
    sg = subgroups[sg_index]
    items = _items(sg)
    name = _clean(sg.get("name"))
    is_acc_subgroup = name.upper() == OTHER_CONDUIT_ACCESSORIES
    if name:
        fill = _theme(*ACCESSORY_HEADING_FILL) if is_acc_subgroup else _theme(*SUBGROUP_FILL)
        sheet.heading(
            _subheading(name, sg.get("note")), fill=fill, center=False, numbered=not is_acc_subgroup
        )
    if is_acc_subgroup:
        for it in items:
            sheet.accessory(_item_description(it), boq_quantity(it), _clean(it.get("unit")) or "EA")
        sheet.blank()
        return

    # The spec may already carry the anchors block as its own subgroup.
    next_is_acc = any(
        _clean(s.get("name")).upper() == OTHER_CONDUIT_ACCESSORIES and _items(s)
        for s in subgroups[sg_index + 1 :]
    )
    headed_block = name.upper().startswith("CONDUITS")
    straps = 0
    # Index in `rows` just after the last EMT/RGS run's fitting rows: an
    # unheaded anchors block (LOW VOLTAGE WIRING and the like) goes there.
    last_steel_fitting_end: int | None = None
    in_steel_run = False
    rows: list[tuple[str, str, float, str]] = []  # (kind, desc, qty, unit)
    for idx, it in enumerate(items):
        qty = boq_quantity(it)
        unit = _clean(it.get("unit")).upper() or "EA"
        desc = _item_description(it)
        if it.get("is_accessory"):
            rows.append(("accessory", desc, qty, unit))
            if "STRAP" in desc.upper():
                straps += _ceil(qty)
            if in_steel_run:
                last_steel_fitting_end = len(rows)
            continue
        rows.append(("item", desc, qty, unit))
        parsed = parse_conduit(_clean(it.get("description")))
        in_steel_run = bool(parsed and parsed[1] in ("EMT", "RGS") and unit == "FT")
        if in_steel_run:
            last_steel_fitting_end = len(rows)
        nxt = items[idx + 1] if idx + 1 < len(items) else None
        spec_has_fittings = bool(nxt and nxt.get("is_accessory"))
        if parsed and unit == "FT" and not spec_has_fittings:
            fittings, s = conduit_fittings(_clean(it.get("description")), qty)
            for f_desc, f_qty, f_unit in fittings:
                rows.append(("accessory", f_desc, f_qty, f_unit))
            straps += s
            if in_steel_run:
                last_steel_fitting_end = len(rows)

    block = [] if next_is_acc else other_conduit_accessories(straps)
    if block and not headed_block:
        insert_at = last_steel_fitting_end if last_steel_fitting_end is not None else len(rows)
        rows[insert_at:insert_at] = [("accessory", d, q, u) for d, q, u in block]

    for kind, desc, qty, unit in rows:
        if kind == "item":
            sheet.item(desc, qty, unit)
        else:
            sheet.accessory(desc, qty, unit)
    if block and headed_block:
        sheet.heading(
            OTHER_CONDUIT_ACCESSORIES,
            fill=_theme(*ACCESSORY_HEADING_FILL),
            center=False,
            numbered=False,
        )
        for d, q, u in block:
            sheet.accessory(d, q, u)
    sheet.blank()


def _write_boq_area(
    ws: Worksheet, title: str, area: dict[str, Any], closing_notes: list[str]
) -> None:
    for col, width in BOQ_WIDTHS.items():
        ws.column_dimensions[col].width = width
    t = _text_cell(ws, 2, 2, title)
    _style(t, bold=True, size=12, h="center", wrap=True, border=False)
    ws.row_dimensions[2].height = max(31.5, 16 * _line_count(title, 56))
    for col, label in enumerate(("SR.\nNO.", "DESCRIPTION", "QUANTITY", "UNIT"), start=1):
        c = ws.cell(row=BOQ_HEADER_ROW, column=col, value=label)
        _style(
            c,
            bold=True,
            h="center",
            wrap=col in (1, 4),
            fill=_solid(HEADER_FILL),
            fmt="0" if col == 3 else None,
        )
    ws.row_dimensions[BOQ_HEADER_ROW].height = 30

    sheet = _BoqSheet(ws)
    for section in area.get("sections") or []:
        if not isinstance(section, dict) or not _section_has_items(section):
            continue
        heading = _clean(section.get("heading")).upper()
        sheet.heading(
            _subheading(heading, section.get("note")), fill=_solid(LEVEL0_FILL), center=True
        )
        for group in section.get("groups") or []:
            if not isinstance(group, dict) or not _group_has_items(group):
                continue
            gname = _clean(group.get("name"))
            if gname and gname.upper() != heading:
                sheet.heading(
                    _group_heading(gname, group.get("note")), fill=_group_fill(gname), center=True
                )
            for i, sg in enumerate(group.get("subgroups") or []):
                if isinstance(sg, dict) and _items(sg):
                    _write_boq_subgroup(sheet, group, i)

    notes = [n for n in list(area.get("notes") or []) + list(closing_notes or []) if _clean(n)]
    if notes:
        sheet.row += 1
        for n in notes:
            sheet.note(n)
    ws.print_area = f"A1:E{max(sheet.row - 1, BOQ_HEADER_ROW)}"


def render_boq_xlsx(boq_spec: dict[str, Any]) -> bytes:
    """`BOQ - COUNTS.xlsx` from MAIN_AGENT_SCHEMA['boq']."""
    areas = _areas(boq_spec)
    wb = Workbook()
    wb.active.title = "Sheet1"
    title = _boq_title(boq_spec)
    closing = [n for n in boq_spec.get("closing_notes") or [] if _clean(n)]
    for area, sheet_name in zip(areas, _area_titles(areas, "SHEET"), strict=True):
        ws = wb.create_sheet(sheet_name)
        _write_boq_area(ws, title, area, closing)
    wb.active = 1
    return _workbook_bytes(wb)


# ---------------------------------------------------------------------------
# <PROJECT> - ELECTRICAL ESTIMATE.xlsx
# ---------------------------------------------------------------------------

EST_WIDTHS = {
    "A": 6.453125, "B": 13.1796875, "C": 68.1796875, "D": 9.453125, "E": 10.36328125,
    "F": 10.08984375, "G": 6.81640625, "H": 9.08984375, "I": 9.1796875, "J": 9.453125,
    "K": 10.6328125, "L": 8.453125, "M": 9.7265625, "N": 10.1796875, "O": 8.6328125,
    "P": 9.7265625, "Q": 9.7265625, "R": 10.6328125,
}  # fmt: skip
EST_HEADERS = (
    "SR.\nNO.",
    "DRAWING SHEET#",
    "DESCRIPTION",
    "QUANTITY",
    "WASTAGE",
    "QTY WITH\nWASTAGE",
    "UNIT",
    "UNIT LABOR HOURS",
    "UNIT MATERIAL COST",
    "UNIT LABOR COST",
    "TOTAL LABOR HOURS",
    "TOTAL MATERIAL COST",
    "TOTAL LABOR COST",
    "UNIT ITEM COST",
    "TOTAL ITEM COST",
)
EST_HEADER_ROW = 7
EST_FIRST_ROW = 8
EST_LAST_COL = 15  # O
SUBTOTAL_LABELS = ("SUBTOTAL MATERIAL", "SUBTOTAL LABOR", "SUBTOTAL")  # P, Q, R

RECAP_SHEET = "Bid Recap & Summary"
RECAP_HEADERS = (
    "SR. NO.",
    "DESCRIPTION",
    "MATERIAL COST",
    "LABOR COST",
    "MATERIAL TAX",
    "LABOR TAX",
    "TOTAL COST",
    "OHP (LABOR)",
    "OHP (MATERIAL)",
    "TOTAL PRICE",
)
RECAP_WIDTHS = {
    "A": 7.54, "B": 11.91, "C": 40.91, "D": 12.18, "E": 13.63, "F": 11.45,
    "G": 11.09, "H": 13.36, "I": 12.18, "J": 13.36, "K": 15.0,
}  # fmt: skip
# (number in column B or None, label, bold, size) below the BID SUMMARY title.
RECAP_SUMMARY_ROWS: tuple[tuple[int | None, str, bool, float], ...] = (
    (1, "TOTAL MATERIAL COST", True, 12),
    (None, "MATERIAL SALES TAX", False, 11),
    (2, "TOTAL LABOR COST", True, 12),
    (None, "LABOR TAX", False, 11),
    (3, "TOTAL COST", True, 12),
    (None, "OVERHEAD & PROFIT (LABOR) @", False, 11),
    (None, "OVERHEAD & PROFIT (MATERIAL) @", False, 11),
    (4, "TOTAL COST WITH OVERHEADS + PROFIT", True, 12),
    (None, "BID SECURITY", False, 11),
)
# Copied verbatim from the sent estimates (grammar included) so a generated
# workbook reads exactly like the ones customers already receive.
ESTIMATE_DISCLAIMER = (
    "Terms and Conditions/Disclaimer: The quantities, numbers, scope, or any other information "
    "contained in this document have been compiled to the best of our knowledge. By utilizing or "
    "deriving any of the information contained in this document, you accept the accuracy of the "
    "information, quantities, and costs, and waive any claims for errors in the information. We does "
    "not warranty any of the information contained in this document. The information provided in "
    "this document is for informational purposes only and should be used as a guideline for "
    "quantities and pricing. It is the responsibility of the contractor performing the work to review "
    "quantities, bid documents, and finalize pricing and materials based on current prices."
)


class _EstimateSheet:
    """Writes one detailed estimate sheet row by row, tracking the serial
    formula in column A."""

    def __init__(self, ws: Worksheet) -> None:
        self.ws = ws
        self.row = EST_FIRST_ROW

    def _serial(self, r: int) -> None:
        start = "A$7:A7" if r == EST_FIRST_ROW else f"A${EST_FIRST_ROW}:A{r - 1}"
        c = self.ws.cell(row=r, column=1, value=f'=IF(D{r}<>"",1+MAX({start}),"")')
        _style(c, bold=True, size=12, h="center", v="center", fmt="0")

    def blank(self) -> None:
        self._serial(self.row)
        self.row += 1

    def heading(self, text: str, *, fill: PatternFill, center: bool) -> None:
        r = self.row
        self._serial(r)
        c = _text_cell(self.ws, r, 3, text)
        _style(c, bold=True, h="center" if center else "left", v="center", wrap=True, fill=fill)
        if "\n" in text:
            self.ws.row_dimensions[r].height = 15 * _line_count(text, 62)
        self.row += 1

    def item(self, description: str, quantity: float, wastage: float, unit: str) -> int:
        r = self.row
        self._serial(r)
        ws = self.ws
        c = _text_cell(ws, r, 3, description)
        _style(c, v="center", wrap=True)
        q = quantity if not float(quantity).is_integer() else int(quantity)
        d = ws.cell(row=r, column=4, value=q)
        _style(d, h="center", v="center", wrap=True, fmt=_qty_format(q))
        e = ws.cell(row=r, column=5, value=wastage if wastage else 0)
        _style(e, h="center", v="center", fmt="0%")
        f = ws.cell(row=r, column=6, value=f"=D{r}*(1+E{r})")
        # Whole quantities keep the sent files' whole-number display; a
        # fractional one must not round to 0 in column F either.
        _style(f, h="center", v="center", fmt=_qty_format(q))
        g = _text_cell(ws, r, 7, _clean(unit).upper())
        _style(g, h="center", v="center")
        for col in range(8, EST_LAST_COL + 1):  # pricing cells: present, empty
            _style(ws.cell(row=r, column=col), h="center", v="center", wrap=True)
        lines = _line_count(description, 62)
        if lines > 1:
            ws.row_dimensions[r].height = 15 * lines
        self.row += 1
        return r

    def sheet_refs(self, first: int, last: int, refs: list[str]) -> None:
        """The DRAWING SHEET# cell, merged over the subgroup's rows."""
        if not refs:
            return
        c = _text_cell(self.ws, first, 2, ", ".join(refs))
        _style(c, bold=True, size=12, h="center", v="center", wrap=True, fmt="0")
        if last > first:
            self.ws.merge_cells(start_row=first, start_column=2, end_row=last, end_column=2)

    def subtotal(self) -> None:
        """The SUBTOTAL label row (P/Q/R) and the empty value row under it.
        The value cells stay blank: AI Estimating never prices."""
        r = self.row
        self._serial(r)
        c = self.ws.cell(row=r, column=3, value="")
        _style(c, v="center", wrap=True)
        for i, label in enumerate(SUBTOTAL_LABELS):
            cell = self.ws.cell(row=r, column=16 + i, value=label)
            _style(cell, h="center", v="center", wrap=True, fill=_solid(SUBTOTAL_FILL))
        self._serial(r + 1)
        for i in range(len(SUBTOTAL_LABELS)):
            _style(
                self.ws.cell(row=r + 1, column=16 + i),
                bold=True,
                h="center",
                v="center",
                fill=_solid(SUBTOTAL_FILL),
            )
        self.row += 2

    def allowance(self, text: str) -> None:
        """One allowance row: text in C, LS in G, the amount left blank."""
        r = self.row
        self._serial(r)
        c = _text_cell(self.ws, r, 3, text)
        _style(c, wrap=True)
        _style(self.ws.cell(row=r, column=4), h="center", v="center", wrap=True)
        g = self.ws.cell(row=r, column=7, value="LS")
        _style(g, h="center", v="center", wrap=True)
        self.row += 1

    def note(self, text: str, *, first: bool) -> None:
        """One note row merged over B:O; the first one carries the yellow
        NOTE: cell in column A."""
        r = self.row
        if first:
            a = self.ws.cell(row=r, column=1, value="NOTE:")
            _style(a, bold=True, size=12, h="center", v="center", fill=_solid(NOTE_FILL))
        body = _clean(text)
        b = _text_cell(self.ws, r, 2, body)
        _style(b, v="center", wrap=True)
        self.ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=EST_LAST_COL)
        if body:
            self.ws.row_dimensions[r].height = max(15.75, 15 * _line_count(body, 150))
        self.row += 1


def _write_estimate_header(
    ws: Worksheet, estimate_spec: dict[str, Any], date: datetime.date
) -> None:
    for col, width in EST_WIDTHS.items():
        ws.column_dimensions[col].width = width
    project = _clean(estimate_spec.get("project_name"))
    address = _clean(estimate_spec.get("address"))
    for r, label, value in (
        (2, "PROJECT", project),
        (3, "ESTIMATE", "ELECTRICAL ESTIMATE"),
        (4, "ADDRESS", address),
    ):
        lc = ws.cell(row=r, column=4, value=label)
        _style(lc, bold=True, size=12, h="left", v="center", fmt="0")
        vc = _text_cell(ws, r, 5, value)
        _style(
            vc,
            bold=True,
            size=12,
            h="center",
            v="center",
            wrap=True,
            color=PROJECT_RED,
            border=r != 2,
        )
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    ws.row_dimensions[2].height = max(16.5, 16 * _line_count(project, 40))
    ws.row_dimensions[4].height = 36.75
    rate = ws.cell(row=4, column=12, value="LABOR RATE")
    _style(rate, bold=True, size=14, h="center", v="center", fill=_solid(HEADER_FILL))
    ws.merge_cells(start_row=4, start_column=12, end_row=4, end_column=13)
    _style(
        ws.cell(row=4, column=14),
        bold=True,
        size=14,
        h="center",
        v="center",
        fill=_solid(HEADER_FILL),
    )
    a5 = ws.cell(row=5, column=1, value="DATE: ")
    _style(a5, bold=True, size=12, h="left")
    b5 = ws.cell(row=5, column=2, value=datetime.datetime(date.year, date.month, date.day))
    _style(
        b5,
        bold=True,
        size=12,
        h="left",
        v="center",
        wrap=True,
        fill=_theme(*DATE_FILL),
        fmt="mm-dd-yy",
    )
    for col, label in enumerate(EST_HEADERS, start=1):
        c = ws.cell(row=EST_HEADER_ROW, column=col, value=label)
        _style(
            c,
            bold=True,
            h="center",
            v="center",
            wrap=col != 3,
            fill=_solid(HEADER_FILL),
            fmt="0" if col == 4 else None,
        )
    ws.row_dimensions[EST_HEADER_ROW].height = 43.2


def _write_estimate_area(
    ws: Worksheet,
    area: dict[str, Any],
    estimate_spec: dict[str, Any],
    date: datetime.date,
    area_notes: list[str],
) -> list[str]:
    """One detailed estimate sheet. Returns the recap labels in order
    ('DEMOLITION' and the NEW WORK group names that have items)."""
    _write_estimate_header(ws, estimate_spec, date)
    sheet = _EstimateSheet(ws)
    labels: list[str] = []
    for section in area.get("sections") or []:
        if not isinstance(section, dict) or not _section_has_items(section, estimate=True):
            continue
        heading = _clean(section.get("heading")).upper()
        sheet.heading(
            _subheading(heading, section.get("note")), fill=_solid(LEVEL0_FILL), center=True
        )
        demolition = heading == "DEMOLITION"
        for group in section.get("groups") or []:
            if not isinstance(group, dict) or not _group_has_items(group, estimate=True):
                continue
            gname = _clean(group.get("name"))
            if gname and gname.upper() != heading:
                sheet.heading(
                    _group_heading(gname, group.get("note")), fill=_group_fill(gname), center=True
                )
            for sg in group.get("subgroups") or []:
                if not isinstance(sg, dict):
                    continue
                sname = _clean(sg.get("name"))
                items = _estimate_items(sg)  # fittings live in the BOQ only
                if not items:
                    continue
                if sname:
                    sheet.heading(
                        _subheading(sname, sg.get("note")),
                        fill=_theme(*SUBGROUP_FILL),
                        center=False,
                    )
                first = last = 0
                refs: list[str] = []
                for it in items:
                    q = round(
                        _num(
                            it.get("quantity"),
                            f"Estimate item '{_clean(it.get('description'))[:60]}'",
                        ),
                        2,
                    )
                    r = sheet.item(
                        _item_description(it),
                        q,
                        estimate_wastage(it),
                        _clean(it.get("unit")) or "EA",
                    )
                    first = first or r
                    last = r
                    for ref in it.get("sheet_refs") or []:
                        ref_s = _clean(ref)
                        if ref_s and ref_s not in refs:
                            refs.append(ref_s)
                sheet.sheet_refs(first, last, refs)
                sheet.blank()
            if not demolition:
                sheet.subtotal()
                labels.append(gname or heading)
        if demolition:
            sheet.subtotal()
            labels.append("DEMOLITION")

    allowances = []
    if estimate_spec.get("include_standard_allowances", True):
        allowances.extend(ESTIMATE_ALLOWANCES)
    allowances.extend(_clean(a) for a in estimate_spec.get("extra_allowances") or [] if _clean(a))
    if allowances:
        sheet.heading("ALLOWANCES", fill=_theme(*SUBGROUP_FILL), center=False)
        for a in allowances:
            sheet.allowance(a)
        labels.append("ALLOWANCES")

    sheet.row += 1
    notes = [n for n in list(estimate_spec.get("general_notes") or []) + area_notes if _clean(n)]
    if not notes:
        sheet.note("", first=True)
    for i, n in enumerate(notes):
        sheet.note(n, first=i == 0)
    ws.print_area = f"A1:R{sheet.row - 1}"
    return labels


def _write_recap(ws: Worksheet, labels: list[str]) -> None:
    for col, width in RECAP_WIDTHS.items():
        ws.column_dimensions[col].width = width
    t = ws.cell(row=2, column=2, value="BID RECAP")
    _style(t, bold=True, size=16, h="center", v="center", fill=_solid(RECAP_TITLE_FILL))
    ws.merge_cells("B2:K2")
    ws.row_dimensions[2].height = 30
    for i, label in enumerate(RECAP_HEADERS):
        c = ws.cell(row=3, column=2 + i, value=label)
        _style(c, bold=True, h="center", v="center", wrap=i != 2)
    ws.row_dimensions[3].height = 27.75

    r = 4
    for label in labels:
        b = ws.cell(row=r, column=2, value=f"=1+MAX($B$3:B{r - 1})")
        _style(b, h="center", v="center")
        c = _text_cell(ws, r, 3, label)
        low_voltage = label.upper().startswith("LOW VOLTAGE WORK")
        _style(
            c,
            h="left",
            v="center",
            wrap=True,
            fill=_theme(*LOW_VOLTAGE_FILL) if low_voltage else None,
        )
        for col in range(4, 12):  # cost cells: present, empty
            _style(ws.cell(row=r, column=col), h="center", v="center", wrap=True)
        ws.row_dimensions[r].height = 25.2
        r += 1

    totals = r + 1
    tc = ws.cell(row=totals, column=3, value="TOTALS")
    _style(tc, bold=True, size=14, h="center", v="center", wrap=True)
    for col in range(4, 12):
        _style(
            ws.cell(row=totals, column=col), bold=True, size=14, h="center", v="center", wrap=True
        )

    summary = totals + 3
    s = ws.cell(row=summary, column=2, value="BID SUMMARY")
    _style(s, bold=True, size=14, h="center", v="center", fill=_solid(RECAP_TITLE_FILL))
    ws.merge_cells(start_row=summary, start_column=2, end_row=summary, end_column=5)
    ws.row_dimensions[summary].height = 30
    r = summary + 1
    for number, label, bold, size in RECAP_SUMMARY_ROWS:
        if number is not None:
            _style(ws.cell(row=r, column=2, value=number), h="center", v="center")
        _style(
            ws.cell(row=r, column=3, value=label),
            bold=bold,
            size=size,
            h="left" if bold else None,
            v="center",
        )
        _style(ws.cell(row=r, column=5), bold=bold, size=size, v="center")
        ws.row_dimensions[r].height = 25.2
        r += 1
    base = r + 1
    _style(ws.cell(row=base, column=2, value=5), h="center", v="center")
    _style(
        ws.cell(row=base, column=3, value="BASE BID PRICE"),
        bold=True,
        size=14,
        h="left",
        v="center",
    )
    _style(ws.cell(row=base, column=5), bold=True, size=14, v="center", fill=_solid(SUBTOTAL_FILL))
    ws.row_dimensions[base].height = 25.2
    disclaimer = base + 2
    d = ws.cell(row=disclaimer, column=3, value=ESTIMATE_DISCLAIMER)
    _style(d, v="center", wrap=True, border=False)
    ws.row_dimensions[disclaimer].height = 234


def _area_key(name: Any) -> str:
    """Area names compared loosely: upper case, punctuation and spacing
    folded ('Silverado Ranch' == 'SILVERADO-RANCH')."""
    return re.sub(r"[^A-Z0-9]+", " ", _clean(name).upper()).strip()


def _route_area_notes(
    estimate_spec: dict[str, Any], areas: list[dict[str, Any]], titles: list[str]
) -> list[list[str]]:
    """The estimate's area_notes, per area sheet. One area takes every note
    whatever area name the model wrote. With several areas a note goes to
    the area whose name (or sheet title) matches loosely; a note that
    matches none is never dropped: it goes on every sheet prefixed with its
    area name, and a warning is logged."""
    per_sheet: list[list[str]] = [[] for _ in areas]
    keys = [
        {_area_key(a.get("area")), _area_key(t)} - {""} for a, t in zip(areas, titles, strict=True)
    ]
    unmatched: list[str] = []
    for an in estimate_spec.get("area_notes") or []:
        if not isinstance(an, dict):
            continue
        texts = [_clean(n) for n in an.get("notes") or [] if _clean(n)]
        if not texts:
            continue
        if len(areas) == 1:
            per_sheet[0].extend(texts)
            continue
        key = _area_key(an.get("area"))
        idx = next((i for i, k in enumerate(keys) if key and key in k), None)
        if idx is not None:
            per_sheet[idx].extend(texts)
            continue
        label = _clean(an.get("area"))
        log.warning(
            "AI Estimating estimate: area note for %r matches no BOQ area; added to every sheet",
            label[:80],
        )
        unmatched.extend(f"{label}: {t}" if label else t for t in texts)
    for notes in per_sheet:
        notes.extend(unmatched)
    return per_sheet


def _multi_area_recap(area_labels: list[list[str]]) -> list[str]:
    """Recap rows for a multi-area estimate: every group label in first-seen
    order, each row covering that group across all area sheets, then
    ALLOWANCES. Group-level rows keep the WIRING and OTHER ITEMS labels the
    General Material extraction reads (general_material.py), which per-area
    rows would hide."""
    out: list[str] = []
    seen: set[str] = set()
    allowances = False
    for labels in area_labels:
        for label in labels:
            key = label.upper()
            if key == "ALLOWANCES":
                allowances = True
            elif key not in seen:
                seen.add(key)
                out.append(label)
    if allowances:
        out.append("ALLOWANCES")
    return out


def render_estimate_xlsx(
    boq_spec: dict[str, Any], estimate_spec: dict[str, Any], *, date: datetime.date
) -> bytes:
    """`<PROJECT> - ELECTRICAL ESTIMATE.xlsx` from the BOQ spec plus
    MAIN_AGENT_SCHEMA['estimate']. Pricing columns are present and empty.
    The Bid Recap lists one row per group label (DEMOLITION, WIRING, OTHER
    ITEMS, ...) and ALLOWANCES; with several areas each row covers that group
    on every area sheet."""
    areas = _areas(boq_spec)
    if not isinstance(estimate_spec, dict):
        raise ValueError("The estimate spec must be an object.")
    wb = Workbook()
    wb.active.title = "Sheet1"
    titles = _area_titles(areas, "Detailed Estimate")
    notes_per_sheet = _route_area_notes(estimate_spec, areas, titles)
    area_labels: list[list[str]] = []
    for area, title, extra in zip(areas, titles, notes_per_sheet, strict=True):
        ws = wb.create_sheet(title)
        area_labels.append(_write_estimate_area(ws, area, estimate_spec, date, extra))
    recap_labels = area_labels[0] if len(areas) == 1 else _multi_area_recap(area_labels)
    _write_recap(wb.create_sheet(RECAP_SHEET), recap_labels)
    wb.active = 1
    return _workbook_bytes(wb)


# ---------------------------------------------------------------------------
# Trench / Sawcut markup PDFs
# ---------------------------------------------------------------------------

MARKUP_LINE_WIDTH = 7.0  # the sent markups use 7 pt lines on 42x30 sheets
MARKUP_FONT_SIZE = 36.0  # the sent markups label at 36 pt
_MARKUP_RED = (1.0, 0.0, 0.0)
_MARKUP_LABEL_FILL = (1.0, 1.0, 0.0)


def _fraction(v: Any) -> float:
    try:
        f = float(v)
    except (TypeError, ValueError):
        raise ValueError("Markup points must be numbers between 0 and 1.") from None
    if math.isnan(f):
        raise ValueError("Markup points must be numbers between 0 and 1.")
    return min(1.0, max(0.0, f))


def _to_user_space(
    fx: float, fy: float, box: tuple[float, float, float, float], rotation: int
) -> tuple[float, float]:
    """Page fractions of the page AS DISPLAYED (origin top-left, the way the
    rendered PNG the model saw is oriented) to PDF user space, honoring the
    page /Rotate. `box` is the crop box (x0, y0, x1, y1)."""
    x0, y0, x1, y1 = box
    w, h = x1 - x0, y1 - y0
    if rotation == 90:
        u, v = fy * w, fx * h
    elif rotation == 180:
        u, v = (1 - fx) * w, fy * h
    elif rotation == 270:
        u, v = (1 - fy) * w, (1 - fx) * h
    else:
        u, v = fx * w, (1 - fy) * h
    return x0 + u, y0 + v


def _pdf_str(text: str) -> str:
    """Latin-1 safe text escaped for a PDF literal string in a content stream."""
    s = _clean(text).encode("latin-1", "replace").decode("latin-1")
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


# Helvetica advance widths (AFM, 1/1000 em) for ASCII 32..126, so the label
# box fits its text without embedding a font.
_HELV_WIDTHS = (
    278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556,
    1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778,
    667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556,
    333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556,
    556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584,
)  # fmt: skip


def _text_width(text: str, size: float) -> float:
    """Helvetica width of `text` at `size` points."""
    total = 0
    for ch in text:
        o = ord(ch)
        total += _HELV_WIDTHS[o - 32] if 32 <= o <= 126 else 556
    return total * size / 1000


def _rgb(c: tuple[float, float, float]) -> str:
    """A PDF color operand string such as '1 0 0'."""
    return " ".join(f"{v:g}" for v in c)


def _indirect(writer, obj):
    """An indirect reference to `obj` inside `writer`. Appearance streams and
    fonts must be indirect objects, and pypdf 6 only offers that as the
    private PdfWriter._add_object (what its own add_annotation uses), so a
    public add_object wins when a later pypdf ships one. Pin pypdf below 7 so
    an upgrade cannot change this silently; the markup tests fail loudly if
    it does."""
    add = getattr(writer, "add_object", None) or getattr(writer, "_add_object", None)
    if add is None:
        raise ValueError("The markup PDF could not be built with the installed PDF library.")
    return add(obj)


def _markup_font(writer):
    """One Helvetica font object for the whole document, registered as /Helv
    in the AcroForm /DR so the FreeText /DA resolves when a viewer redraws a
    label a user edited."""
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject

    font_ref = _indirect(
        writer,
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
                NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
            }
        ),
    )
    root = writer.root_object
    acro = root.get("/AcroForm")
    if acro is None:
        acro = DictionaryObject(
            {
                NameObject("/Fields"): ArrayObject(),
                NameObject("/DA"): TextStringObject("/Helv 0 Tf 0 g"),
            }
        )
        root[NameObject("/AcroForm")] = acro
    else:
        acro = acro.get_object()
    dr = acro.get("/DR")
    if dr is None:
        dr = DictionaryObject()
        acro[NameObject("/DR")] = dr
    else:
        dr = dr.get_object()
    fonts = dr.get("/Font")
    if fonts is None:
        fonts = DictionaryObject()
        dr[NameObject("/Font")] = fonts
    else:
        fonts = fonts.get_object()
    if "/Helv" not in fonts:
        fonts[NameObject("/Helv")] = font_ref
    return font_ref


def _markup_line_annotation(
    writer, p1: tuple[float, float], p2: tuple[float, float], width: float, contents: str
):
    """One red /Line annotation from p1 to p2 (user space) with its own
    appearance stream."""
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        FloatObject,
        NameObject,
        NumberObject,
        TextStringObject,
    )

    pad = width / 2 + 1
    rect = [
        min(p1[0], p2[0]) - pad,
        min(p1[1], p2[1]) - pad,
        max(p1[0], p2[0]) + pad,
        max(p1[1], p2[1]) + pad,
    ]
    ap = DecodedStreamObject()
    ap.set_data(
        f"q {_rgb(_MARKUP_RED)} RG {width:g} w 1 J 1 j {p1[0]:.3f} {p1[1]:.3f} m {p2[0]:.3f} {p2[1]:.3f} l S Q".encode()
    )
    ap.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Form"),
            NameObject("/BBox"): ArrayObject([FloatObject(v) for v in rect]),
        }
    )
    annot = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/Line"),
            NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]),
            NameObject("/L"): ArrayObject([FloatObject(v) for v in (*p1, *p2)]),
            NameObject("/C"): ArrayObject([FloatObject(v) for v in _MARKUP_RED]),
            NameObject("/BS"): DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Border"),
                    NameObject("/S"): NameObject("/S"),
                    NameObject("/W"): FloatObject(width),
                }
            ),
            NameObject("/Border"): ArrayObject(
                [NumberObject(0), NumberObject(0), FloatObject(width)]
            ),
            NameObject("/LE"): ArrayObject([NameObject("/None"), NameObject("/None")]),
            NameObject("/F"): NumberObject(4),
            NameObject("/CA"): FloatObject(1),
            NameObject("/AP"): DictionaryObject({NameObject("/N"): _indirect(writer, ap)}),
        }
    )
    if contents:
        annot[NameObject("/Contents")] = TextStringObject(contents)
    return annot


MARKUP_MIN_FONT_SIZE = 10.0  # a long label shrinks to this before it wraps
_LABEL_MAX_WIDTH = 0.9  # share of the displayed page width a label box may use


def _wrap_pdf_line(text: str, size: float, max_w: float) -> list[str]:
    """Word-wrap one line of Helvetica text to `max_w` points; a single word
    wider than that is broken by characters."""
    lines: list[str] = []
    line = ""
    for word in text.split():
        trial = f"{line} {word}" if line else word
        if _text_width(trial, size) <= max_w:
            line = trial
            continue
        if line:
            lines.append(line)
        line = word
        while len(line) > 1 and _text_width(line, size) > max_w:
            cut = len(line) - 1
            while cut > 1 and _text_width(line[:cut], size) > max_w:
                cut -= 1
            lines.append(line[:cut])
            line = line[cut:]
    if line:
        lines.append(line)
    return lines


def _label_layout(
    label: str, notes: str, size: float, max_w: float
) -> tuple[list[str], float, list[str], float, float, float, float]:
    """(label lines, label size, note lines, note size, pad, box width, box
    height) for a label box at most `max_w` points wide as displayed. Each
    newline in the label or notes starts its own line; a box wider than the
    limit first shrinks its font (not below MARKUP_MIN_FONT_SIZE), then wraps."""
    label_raw = [ln.strip() for ln in _clean(label).split("\n") if ln.strip()]
    note_raw = [ln.strip() for ln in _clean(notes).split("\n") if ln.strip()]
    max_w = max(max_w, 4 * MARKUP_MIN_FONT_SIZE)

    def natural(s: float) -> float:
        ns = max(8.0, s * 0.5)
        widths = [_text_width(t, s) for t in label_raw] + [_text_width(t, ns) for t in note_raw]
        return max(widths or [0.0]) + 2 * s * 0.25

    width = natural(size)
    if width > max_w:
        size = max(MARKUP_MIN_FONT_SIZE, size * max_w / width)
    note_size = max(8.0, size * 0.5)
    pad = size * 0.25
    text_w = max_w - 2 * pad
    label_lines = [w for t in label_raw for w in _wrap_pdf_line(t, size, text_w)]
    note_lines = [w for t in note_raw for w in _wrap_pdf_line(t, note_size, text_w)]
    widths = [_text_width(t, size) for t in label_lines] + [
        _text_width(t, note_size) for t in note_lines
    ]
    bw = max(widths or [0.0]) + 2 * pad
    bh = (
        2 * pad
        + len(label_lines) * size * 1.15
        + ((size * 0.3 + len(note_lines) * note_size * 1.25) if note_lines else 0)
    )
    return label_lines, size, note_lines, note_size, pad, bw, bh


def _markup_label_annotation(
    writer,
    font_ref,
    anchor: tuple[float, float],
    box: tuple[float, float, float, float],
    rotation: int,
    label: str,
    notes: str,
    size: float,
):
    """The yellow /FreeText label box with red text just above `anchor` as
    displayed, counter-rotated to read upright and kept inside the page."""
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        FloatObject,
        NameObject,
        NumberObject,
        TextStringObject,
    )

    x0, y0, x1, y1 = box
    display_w = (y1 - y0) if rotation in (90, 270) else (x1 - x0)
    label_lines, size, note_lines, note_size, pad, bw, bh = _label_layout(
        label, notes, size, display_w * _LABEL_MAX_WIDTH
    )

    ops = [f"q {_rgb(_MARKUP_LABEL_FILL)} rg 0 0 {bw:.2f} {bh:.2f} re f Q"]
    baseline = bh - pad - size * 0.9
    for i, ln in enumerate(label_lines):
        if i:
            baseline -= size * 1.15
        ops.append(
            f"BT /Helv {size:g} Tf {_rgb(_MARKUP_RED)} rg {pad:.2f} {baseline:.2f} Td ({_pdf_str(ln)}) Tj ET"
        )
    baseline -= size * 0.3
    for ln in note_lines:
        baseline -= note_size * 1.25
        ops.append(
            f"BT /Helv {note_size:g} Tf {_rgb(_MARKUP_RED)} rg {pad:.2f} {baseline:.2f} Td ({_pdf_str(ln)}) Tj ET"
        )
    ap = DecodedStreamObject()
    ap.set_data(" ".join(ops).encode("latin-1"))
    # Counter-rotate so the label reads upright on a rotated sheet.
    rad = math.radians(rotation)
    cos, sin = round(math.cos(rad), 6), round(math.sin(rad), 6)
    ap.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Form"),
            NameObject("/BBox"): ArrayObject(
                [FloatObject(0), FloatObject(0), FloatObject(bw), FloatObject(bh)]
            ),
            NameObject("/Matrix"): ArrayObject(
                [FloatObject(v) for v in (cos, sin, -sin, cos, 0, 0)]
            ),
            NameObject("/Resources"): DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/Helv"): font_ref})}
            ),
        }
    )
    # Rect in user space: the box is bw x bh as displayed; swap on 90/270.
    rw, rh = (bh, bw) if rotation in (90, 270) else (bw, bh)
    gap = size * 0.5
    ax, ay = anchor
    # Above the first point as displayed, clamped inside the page.
    if rotation == 0:
        rx, ry = ax, ay + gap
    elif rotation == 180:
        rx, ry = ax - rw, ay - gap - rh
    elif rotation == 90:
        rx, ry = ax - gap - rw, ay
    else:
        rx, ry = ax + gap, ay - rh
    rx = min(max(rx, x0), max(x0, x1 - rw))
    ry = min(max(ry, y0), max(y0, y1 - rh))
    rect = [rx, ry, rx + rw, ry + rh]
    contents = "\n".join(label_lines + note_lines)
    annot = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/FreeText"),
            NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]),
            NameObject("/Contents"): TextStringObject(contents),
            NameObject("/DA"): TextStringObject(f"/Helv {size:g} Tf {_rgb(_MARKUP_RED)} rg"),
            NameObject("/DS"): TextStringObject(f"font: Helvetica {size:g}pt; color: #FF0000"),
            NameObject("/C"): ArrayObject([FloatObject(v) for v in _MARKUP_LABEL_FILL]),
            NameObject("/BS"): DictionaryObject({NameObject("/W"): NumberObject(0)}),
            NameObject("/Border"): ArrayObject([NumberObject(0), NumberObject(0), NumberObject(0)]),
            NameObject("/F"): NumberObject(4),
            NameObject("/AP"): DictionaryObject({NameObject("/N"): _indirect(writer, ap)}),
        }
    )
    return annot


MARKUP_KINDS: dict[str, str] = {
    "trench": "trench",
    "sawcut": "sawcut",
    "trench_markup": "trench",
    "sawcut_markup": "sawcut",
}
_SINGLE_SOURCE = "__single__"
MarkupSources = bytes | Mapping[Any, bytes | tuple[bytes, int]]


def _markup_kind(kind: Any) -> str:
    """'trench' or 'sawcut' from a sheet kind or a deliverable kind."""
    k = MARKUP_KINDS.get(_clean(kind).lower())
    if k is None:
        raise ValueError("A markup kind must be trench or sawcut.")
    return k


def _sheet_doc_key(sheet: dict[str, Any]) -> Any:
    """The document a markup sheet points at: its document_id (added by the
    app beside every doc_index) when present, else its doc_index, else None."""
    doc_id = sheet.get("document_id")
    if doc_id not in (None, ""):
        return str(doc_id)
    doc_index = sheet.get("doc_index")
    if doc_index in (None, ""):
        return None
    try:
        return int(doc_index)
    except (TypeError, ValueError):
        raise ValueError("A markup sheet has an invalid document number.") from None


def _sheets_of_kind(sheets: Any, kind: str) -> list[dict[str, Any]]:
    """The sheets of one markup kind. MARKUPS_SPEC mixes trench and sawcut
    sheets in one list; each file carries only its own kind."""
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("There are no markup sheets to render.")
    out = []
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise ValueError("Every markup sheet must be an object.")
        if _markup_kind(sheet.get("kind")) == kind:
            out.append(sheet)
    if not out:
        raise ValueError(f"There are no {kind} markup sheets to render.")
    return out


def markup_documents(sheets: list[dict[str, Any]], *, kind: str) -> list[Any]:
    """The document keys (document_id, else doc_index) the `kind` markup
    needs, in first-reference order, so the caller downloads only those PDFs
    and passes them to render_markup_pdf as `sources`."""
    k = _markup_kind(kind)
    out: list[Any] = []
    for sheet in _sheets_of_kind(sheets, k):
        key = _sheet_doc_key(sheet)
        if key is not None and key not in out:
            out.append(key)
    return out


def _resolve_source(sources: MarkupSources, sheet: dict[str, Any]) -> tuple[Any, bytes, int]:
    """(source key, pdf bytes, first page) for one sheet. `first page` is the
    page of the PDF that is page 1 of the sheet's document."""
    if isinstance(sources, (bytes, bytearray)):
        return _SINGLE_SOURCE, bytes(sources), 1
    if not isinstance(sources, Mapping):
        raise ValueError("The markup needs the source drawing PDFs.")
    candidates = [sheet.get("document_id"), _sheet_doc_key(sheet), sheet.get("doc_index")]
    for key in candidates:
        if key in (None, ""):
            continue
        for k in (key, str(key)):
            if k in sources:
                value = sources[k]
                if isinstance(value, tuple):
                    data, first = value
                else:
                    data, first = value, 1
                try:
                    first_page = max(1, int(first))
                except (TypeError, ValueError):
                    raise ValueError("A markup source has an invalid first page.") from None
                if not isinstance(data, (bytes, bytearray)):
                    raise ValueError("A markup source drawing is not a PDF file.")
                return k, bytes(data), first_page
    raise ValueError(
        f"A {sheet.get('kind') or 'markup'} sheet references document "
        f"{_clean(sheet.get('doc_index')) or '?'}, which was not provided."
    )


def render_markup_pdf(
    sources: MarkupSources,
    sheets: list[dict[str, Any]],
    *,
    kind: str,
    line_width: float = MARKUP_LINE_WIDTH,
    font_size: float = MARKUP_FONT_SIZE,
) -> bytes:
    """`Trench Markup.pdf` (kind 'trench' or 'trench_markup') or `Sawcut
    Markup.pdf` (kind 'sawcut' or 'sawcut_markup'): only the sheets of that
    kind are drawn, so the whole MARKUPS_SPEC can be passed for both files.
    Copies the referenced pages (first-reference order; sheets on the same
    page share it) and adds one red /Line per polyline segment plus one
    /FreeText label per sheet near its first line.

    `sources` maps each document key to that document's PDF bytes: the key
    is the sheet's `document_id` (what the app adds beside every doc_index)
    or its `doc_index`; markup_documents() lists the keys needed. A value may
    be `(pdf_bytes, first_page)` when the bytes are a whole source file and
    the document starts at `first_page` of it (plans_reader_documents
    page_start); a splitter segment's own PDF needs no offset. Passing bare
    bytes means every sheet comes from that one document, and sheets that
    name different documents are refused.

    Each sheet is `{kind, doc_index, document_id?, page (1-based in its
    document), label, lines: [{points: [[fx, fy], ...], length_ft, label}],
    notes}`; points are fractions of the page as displayed, origin top-left."""
    from pypdf import PdfReader, PdfWriter

    k = _markup_kind(kind)
    chosen = _sheets_of_kind(sheets, k)
    if isinstance(sources, (bytes, bytearray)):
        docs = {_sheet_doc_key(s) for s in chosen} - {None}
        if len(docs) > 1:
            raise ValueError(
                "The markup sheets come from more than one document, so each document's PDF is needed."
            )

    # Validate everything before pypdf touches the third-party file, so the
    # only errors left after this pass are malformed-PDF ones.
    plans = []
    for sheet in chosen:
        try:
            page_no = int(sheet.get("page"))
        except (TypeError, ValueError):
            raise ValueError("Every markup sheet needs a page number.") from None
        if page_no < 1:
            raise ValueError("Markup page numbers start at 1.")
        src_key, data, first_page = _resolve_source(sources, sheet)
        lines = []
        for line in sheet.get("lines") or []:
            line = line if isinstance(line, dict) else {}
            pts = []
            for p in line.get("points") or []:
                if not isinstance(p, (list, tuple)) or len(p) < 2:
                    raise ValueError("Markup points must be [x, y] pairs.")
                pts.append((_fraction(p[0]), _fraction(p[1])))
            if len(pts) < 2:
                continue
            parts = [_clean(line.get("label"))]
            if line.get("length_ft") is not None:
                try:
                    parts.append(f"{float(line['length_ft']):g} FT")
                except (TypeError, ValueError):
                    pass
            lines.append((pts, " ".join(p for p in parts if p)))
        label = _clean(sheet.get("label"))
        notes = _clean(sheet.get("notes"))
        plans.append(
            {
                "src": src_key,
                "data": data,
                "abs_page": first_page - 1 + page_no,
                "page": page_no,
                "lines": lines,
                # No label: the notes become the label, one line per newline.
                "label": label or notes,
                "notes": notes if label else "",
            }
        )

    readers: dict[Any, Any] = {}
    for plan in plans:
        if plan["src"] in readers:
            continue
        try:
            reader = PdfReader(io.BytesIO(plan["data"]))
            readers[plan["src"]] = (reader, len(reader.pages))
        except Exception:
            raise ValueError("The source drawing could not be read as a PDF.") from None
    for plan in plans:
        count = readers[plan["src"]][1]
        if plan["abs_page"] > count:
            raise ValueError(
                f"Markup page {plan['page']} is outside the source drawing ({count} pages)."
            )

    try:
        writer = PdfWriter()
        font_ref = None
        out_index: dict[tuple[Any, int], int] = {}
        for plan in plans:
            key = (plan["src"], plan["abs_page"])
            if key not in out_index:
                writer.add_page(readers[plan["src"]][0].pages[plan["abs_page"] - 1])
                out_index[key] = len(writer.pages) - 1
        for plan in plans:
            idx = out_index[(plan["src"], plan["abs_page"])]
            page = writer.pages[idx]
            rotation = int(page.rotation or 0) % 360
            cb = page.cropbox
            box = (float(cb.left), float(cb.bottom), float(cb.right), float(cb.top))
            first_point: tuple[float, float] | None = None
            for fpts, contents in plan["lines"]:
                pts = [_to_user_space(fx, fy, box, rotation) for fx, fy in fpts]
                if first_point is None:
                    first_point = pts[0]
                for a, b in zip(pts, pts[1:], strict=False):
                    writer.add_annotation(
                        idx, _markup_line_annotation(writer, a, b, line_width, contents)
                    )
            if plan["label"]:
                if first_point is None:
                    first_point = _to_user_space(0.02, 0.04, box, rotation)
                if font_ref is None:
                    font_ref = _markup_font(writer)
                annot = _markup_label_annotation(
                    writer,
                    font_ref,
                    first_point,
                    box,
                    rotation,
                    plan["label"],
                    plan["notes"],
                    font_size,
                )
                writer.add_annotation(idx, annot)
        buf = io.BytesIO()
        writer.write(buf)
    except Exception:
        log.warning(
            "AI Estimating markup: the source drawing could not be processed", exc_info=True
        )
        raise ValueError("The source drawing could not be processed.") from None
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Marked Plans.pdf
# ---------------------------------------------------------------------------

MARKED_DPI = 200
MARKED_PAGE_PX = (11 * MARKED_DPI, int(8.5 * MARKED_DPI))  # 2200 x 1700, letter landscape
MARKED_SHEET_FRACTION = 0.78
MARKED_ALPHA = 150
MARKED_PALETTE: tuple[tuple[int, int, int], ...] = (
    (0, 229, 255),
    (153, 51, 51),
    (255, 0, 255),
    (153, 153, 51),
    (102, 51, 204),
    (0, 200, 0),
    (230, 0, 0),
    (51, 51, 255),
    (255, 230, 0),
    (0, 153, 153),
    (255, 140, 0),
    (128, 0, 128),
    (0, 102, 204),
    (204, 0, 102),
    (102, 153, 0),
    (255, 102, 153),
)


def _font(size: int):
    """Pillow's default font at `size` px (a fixed bitmap font on very old
    Pillow)."""
    from PIL import ImageFont

    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # very old Pillow without sized default fonts
        return ImageFont.load_default()


def _fit_text(draw, text: str, font, max_w: float) -> str:
    """`text` cut with an ellipsis to fit `max_w` px."""
    if draw.textlength(text, font=font) <= max_w:
        return text
    ell = "..."
    lo, hi = 0, len(text)
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if draw.textlength(text[:mid].rstrip() + ell, font=font) <= max_w:
            lo = mid
        else:
            hi = mid - 1
    return text[:lo].rstrip() + ell


def _wrap2(draw, text: str, font, max_w: float) -> list[str]:
    """Up to two lines, the second truncated with an ellipsis."""
    if draw.textlength(text, font=font) <= max_w:
        return [text]
    words = text.split(" ")
    line = ""
    for i, w in enumerate(words):
        trial = f"{line} {w}".strip()
        if draw.textlength(trial, font=font) > max_w and line:
            return [line, _fit_text(draw, " ".join(words[i:]), font, max_w)]
        line = trial
    return [_fit_text(draw, text, font, max_w)]


def _fmt_qty(q: Any) -> str:
    """A legend quantity with one decimal, as on the sent marked plans."""
    try:
        return f"{float(q):.1f}"
    except (TypeError, ValueError):
        return ""


def _legend_color(index: int) -> tuple[int, int, int]:
    """The palette color of a legend row (0-based)."""
    return MARKED_PALETTE[index % len(MARKED_PALETTE)]


def _ink_on(color: tuple[int, int, int]) -> tuple[int, int, int]:
    """Black or white, whichever reads better on `color`."""
    r, g, b = color
    return (0, 0, 0) if (0.299 * r + 0.587 * g + 0.114 * b) > 150 else (255, 255, 255)


def _centered_text(draw, center: tuple[float, float], text: str, font, fill) -> None:
    """`text` centered on `center`."""
    x0, y0, x1, y1 = draw.textbbox((0, 0), text, font=font)
    cx, cy = center
    draw.text((cx - (x1 - x0) / 2 - x0, cy - (y1 - y0) / 2 - y0), text, font=font, fill=fill)


def _load_sheet_image(png: bytes, box_w: int, box_h: int):
    """The sheet image scaled to fit box_w x box_h, as RGBA. It is resized in
    its own RGB mode first (JPEG sources decode straight at the reduced size
    via draft), so a large render never gets a full-size RGBA copy."""
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(png))
        scale = min(box_w / img.width, box_h / img.height)
        dw, dh = max(1, int(img.width * scale)), max(1, int(img.height * scale))
        img.draft("RGB", (dw, dh))
        img.load()
        if img.mode not in ("RGB", "RGBA"):
            transparent = img.mode in ("LA", "PA", "RGBa") or "transparency" in img.info
            img = img.convert("RGBA" if transparent else "RGB")
        img = img.resize((dw, dh), Image.LANCZOS, reducing_gap=3.0)
        return img.convert("RGBA")
    except Exception:
        raise ValueError("A marked plans sheet image could not be read.") from None


def _draw_legend_row(
    d,
    x0: int,
    x1: int,
    y: int,
    size: int,
    index: int,
    row: dict[str, Any],
    *,
    polyline: bool,
    numbered: bool,
) -> None:
    """One legend row between x0 and x1 starting at y: description (two
    lines at most), quantity and the color swatch, square for polylines and
    round for dots, carrying the row number when the legend is numbered."""
    font = _font(size)
    line_h = int(size * 1.25)
    row_h = int(size * 2.6)
    swatch = int(size * (1.6 if numbered else 1.1))
    color = _legend_color(index)
    qty = f"{_fmt_qty(row.get('quantity'))} {_clean(row.get('unit')).upper()}".strip()
    qty_w = d.textlength(qty, font=font)
    sw_x1 = x1
    sw_x0 = sw_x1 - swatch
    sw_y0 = y + (row_h - swatch) // 2 - 4
    shape = (sw_x0, sw_y0, sw_x1, sw_y0 + swatch)
    if polyline:
        d.rectangle(shape, fill=color)
    else:
        d.ellipse(shape, fill=color)
    if numbered:
        _centered_text(
            d,
            ((sw_x0 + sw_x1) / 2, sw_y0 + swatch / 2),
            str(index + 1),
            _font(max(9, int(swatch * 0.5))),
            _ink_on(color),
        )
    qx = sw_x0 - 10 - qty_w
    d.text((qx, y + (row_h - line_h) // 2 - 4), qty, font=font, fill=(0, 0, 200))
    prefix = _clean(row.get("prefix"))
    desc = _clean(row.get("description"))
    text = f"{prefix}, {desc}" if prefix else desc
    if numbered:
        text = f"{index + 1}. {text}"
    lines = _wrap2(d, text, font, qx - 12 - x0)
    ty = y + (row_h - line_h * len(lines)) // 2 - 4
    for ln in lines:
        d.text((x0, ty), ln, font=font, fill=(0, 0, 0))
        ty += line_h


MARKED_LEGEND_MIN_SIZE = 12
MARKED_LEGEND_MAX_SIZE = 22


def _legend_header(d, x0: int, x1: int, y: int, page: dict[str, Any], title: str) -> int:
    """Title, area and sheet number at the top of a legend box; returns the
    y below the divider."""
    title_font = _font(30)
    sub_font = _font(22)
    d.text((x0, y), title, font=title_font, fill=(32, 33, 89))
    y += 42
    for line in (
        _clean(page.get("area")),
        f"SHEET {_clean(page.get('sheet_number'))}" if page.get("sheet_number") else "",
    ):
        if line:
            d.text((x0, y), _fit_text(d, line, sub_font, x1 - x0), font=sub_font, fill=(0, 0, 0))
            y += 30
    y += 8
    d.line((x0, y, x1, y), fill=(200, 200, 200), width=2)
    return y + 12


def _marked_pages(page: dict[str, Any]) -> list:
    """The letter landscape page(s) for one sheet: the marked sheet with its
    legend box, plus legend continuation pages when the legend does not fit
    at the smallest legend font. Every legend row is shown and every mark
    has a legend entry. Past 16 rows the palette repeats, so rows, dots and
    polylines then also carry the row number."""
    from PIL import Image, ImageDraw

    png = page.get("png_bytes")
    if not png:
        raise ValueError("Every marked plans page needs the rendered sheet image.")
    pw, ph = MARKED_PAGE_PX
    margin = 40
    sheet_box_w = int(pw * MARKED_SHEET_FRACTION) - margin
    sheet_box_h = ph - 2 * margin
    sheet = _load_sheet_image(png, sheet_box_w, sheet_box_h)
    dw, dh = sheet.size
    sx, sy = margin, margin + (sheet_box_h - dh) // 2

    legend = [r for r in page.get("legend") or [] if isinstance(r, dict)]
    marks = [m for m in page.get("marks") or [] if isinstance(m, dict)]
    numbered = len(legend) > len(MARKED_PALETTE)
    polyline_rows: set[int] = set()

    overlay = Image.new("RGBA", (dw, dh), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    radius = max(8, round(0.0105 * dw * (1.35 if numbered else 1)))
    stroke = max(5, round(0.006 * dw))
    tag_font = _font(max(10, int(radius * 1.1)))
    tags: list[tuple[float, float, int]] = []
    for m in marks:
        try:
            li = int(m.get("legend_index"))
        except (TypeError, ValueError):
            continue
        if li < 0 or li >= len(legend):
            log.debug("marked plans: mark with legend_index %s outside the legend", li)
            continue
        color = (*_legend_color(li), MARKED_ALPHA)
        pts = []
        for p in m.get("points") or []:
            if isinstance(p, (list, tuple)) and len(p) >= 2:
                try:
                    pts.append((_fraction(p[0]) * dw, _fraction(p[1]) * dh))
                except ValueError:
                    continue
        if not pts:
            continue
        if m.get("kind") == "polyline" and len(pts) >= 2:
            polyline_rows.add(li)
            od.line(pts, fill=color, width=stroke, joint="curve")
            if numbered:
                tags.append((*pts[0], li))
        else:
            for x, y in pts:
                od.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color)
                if numbered:
                    tags.append((x, y, li))
    for x, y, li in tags:  # numbers drawn opaque on top so they stay legible
        if li in polyline_rows:
            od.ellipse(
                (x - radius, y - radius, x + radius, y + radius),
                fill=(255, 255, 255, 230),
                outline=(*_legend_color(li), 255),
                width=3,
            )
            ink = (0, 0, 0, 255)
        else:
            ink = (*_ink_on(_legend_color(li)), 255)
        _centered_text(od, (x, y), str(li + 1), tag_font, ink)
    sheet = Image.alpha_composite(sheet, overlay)
    overlay.close()

    canvas = Image.new("RGB", (pw, ph), "white")
    canvas.paste(sheet.convert("RGB"), (sx, sy))
    sheet.close()
    d = ImageDraw.Draw(canvas)
    d.rectangle((sx - 1, sy - 1, sx + dw, sy + dh), outline=(160, 160, 160), width=2)

    # Legend box on the right.
    lx0 = int(pw * MARKED_SHEET_FRACTION) + 10
    lx1 = pw - margin
    ly0, ly1 = margin, ph - margin
    d.rectangle((lx0, ly0, lx1, ly1), outline=(32, 33, 89), width=3)
    pad = 16
    y = _legend_header(d, lx0 + pad, lx1 - pad, ly0 + pad, page, "MARKED PLANS")
    avail = ly1 - pad - y

    size = MARKED_LEGEND_MAX_SIZE
    while size > MARKED_LEGEND_MIN_SIZE and int(size * 2.6) * len(legend) > avail:
        size -= 2
    row_h = int(size * 2.6)
    if row_h * len(legend) <= avail:
        on_page = len(legend)
    else:  # keep one row's height for the "continued" line
        on_page = max(0, avail // row_h - 1)
    rows = list(enumerate(legend))
    for i, row in rows[:on_page]:
        _draw_legend_row(
            d, lx0 + pad, lx1 - pad, y, size, i, row,
            polyline=i in polyline_rows, numbered=numbered,
        )  # fmt: skip
        y += row_h
    pages = [canvas]
    rest = rows[on_page:]
    if not rest:
        return pages
    d.text(
        (lx0 + pad, y),
        f"Legend continues on the next page ({len(rest)} more)",
        font=_font(size),
        fill=(90, 90, 90),
    )

    # Continuation pages: two legend columns across the whole page.
    size = MARKED_LEGEND_MAX_SIZE
    row_h = int(size * 2.6)
    gap = 40
    col_w = (pw - 2 * margin - 2 * pad - gap) // 2
    while rest:
        cont = Image.new("RGB", (pw, ph), "white")
        cd = ImageDraw.Draw(cont)
        cd.rectangle((margin, margin, pw - margin, ph - margin), outline=(32, 33, 89), width=3)
        top = _legend_header(
            cd,
            margin + pad,
            pw - margin - pad,
            margin + pad,
            page,
            "MARKED PLANS - LEGEND (CONTINUED)",
        )
        per_col = max(1, (ph - margin - pad - top) // row_h)
        for col in range(2):
            x0 = margin + pad + col * (col_w + gap)
            y = top
            for i, row in rest[:per_col]:
                _draw_legend_row(
                    cd, x0, x0 + col_w, y, size, i, row,
                    polyline=i in polyline_rows, numbered=numbered,
                )  # fmt: skip
                y += row_h
            rest = rest[per_col:]
        pages.append(cont)
    return pages


def render_marked_plans_pdf(pages: list[dict[str, Any]]) -> bytes:
    """`Marked Plans.pdf`: one letter landscape page (11 x 8.5 in at 200 dpi)
    per sheet with the sheet image on the left ~78%, translucent count dots
    or polylines colored per legend row, and the legend box on the right
    (plus continuation pages for a legend too long for the box).

    Each page is `{png_bytes, legend: [{prefix, description, quantity, unit}],
    marks: [{legend_index (0-based), kind: dot|polyline, points}], area,
    sheet_number}`; points are fractions of the sheet image, origin top-left.
    Pages are encoded one sheet at a time and merged, so a 40-sheet set
    never holds 40 full-size canvases in memory."""
    from pypdf import PdfReader, PdfWriter

    if not pages:
        raise ValueError("There are no marked plans pages to render.")
    writer = PdfWriter()
    for page in pages:
        for img in _marked_pages(page):
            buf = io.BytesIO()
            img.save(buf, "PDF", resolution=MARKED_DPI, quality=88)
            img.close()
            buf.seek(0)
            writer.append(PdfReader(buf))
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Deep Brief.docx (+ PDF)
# ---------------------------------------------------------------------------


def _add_bullets(doc, items: list[str], style: str = "List Bullet") -> None:
    cleaned = [_clean(i) for i in items if _clean(i)]
    if not cleaned:
        doc.add_paragraph("None noted.")
        return
    for text in cleaned:
        doc.add_paragraph(text, style=style)


def _add_table(doc, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    if not rows:
        doc.add_paragraph("None noted.")
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = _clean(v)


def render_brief_docx(brief_spec: dict[str, Any], *, project_name: str) -> bytes:
    """`Deep Brief.docx` from MAIN_AGENT_SCHEMA['brief']: one heading per
    section, bullet lists, and tables for systems, risks, sheet summaries and
    spec highlights."""
    from docx import Document
    from docx.shared import Pt

    if not isinstance(brief_spec, dict):
        raise ValueError("The brief spec must be an object.")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    name = _clean(project_name) or "Project"
    doc.add_heading(f"{name} Deep Brief", level=0)
    doc.add_paragraph(
        "Electrical scope only. Quantities and scope, no pricing. Prepared by AI Estimating for estimator review."
    )

    doc.add_heading("Project summary", level=1)
    doc.add_paragraph(_clean(brief_spec.get("project_summary")) or "None noted.")
    doc.add_heading("Electrical scope summary", level=1)
    doc.add_paragraph(_clean(brief_spec.get("electrical_scope_summary")) or "None noted.")

    doc.add_heading("Systems", level=1)
    _add_table(
        doc,
        ("System", "In scope", "Notes"),
        [
            (s.get("system"), _clean(s.get("in_scope")).upper(), s.get("notes"))
            for s in brief_spec.get("systems") or []
            if isinstance(s, dict)
        ],
    )

    doc.add_heading("Assumptions", level=1)
    _add_bullets(doc, list(brief_spec.get("assumptions") or []))
    doc.add_heading("Exclusions", level=1)
    _add_bullets(doc, list(brief_spec.get("exclusions") or []))

    doc.add_heading("Risks", level=1)
    _add_table(
        doc,
        ("Risk", "Severity", "Source"),
        [
            (r.get("risk"), _clean(r.get("severity")).upper(), r.get("source"))
            for r in brief_spec.get("risks") or []
            if isinstance(r, dict)
        ],
    )

    doc.add_heading("RFIs and confirmations", level=1)
    rfis = []
    for q in brief_spec.get("rfis") or []:
        if isinstance(q, dict) and _clean(q.get("question")):
            src = _clean(q.get("source"))
            rfis.append(
                f"{_clean(q.get('question'))} ({src})" if src else _clean(q.get("question"))
            )
    _add_bullets(doc, rfis, style="List Number")

    doc.add_heading("Missing documents", level=1)
    _add_bullets(doc, list(brief_spec.get("missing_documents") or []))
    doc.add_heading("Other trades affecting electrical", level=1)
    _add_bullets(doc, list(brief_spec.get("other_trades_affecting_electrical") or []))

    doc.add_heading("Sheet summaries", level=1)
    _add_table(
        doc,
        ("Sheet", "Title", "Summary"),
        [
            (s.get("sheet_number"), s.get("title"), s.get("summary"))
            for s in brief_spec.get("sheet_summaries") or []
            if isinstance(s, dict)
        ],
    )
    doc.add_heading("Specification highlights", level=1)
    _add_table(
        doc,
        ("Section", "Title", "Summary"),
        [
            (s.get("section"), s.get("title"), s.get("summary"))
            for s in brief_spec.get("spec_highlights") or []
            if isinstance(s, dict)
        ],
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def office_pdf_bytes(content: bytes, filename: str) -> bytes | None:
    """An office deliverable (the BOQ and estimate xlsx, the brief docx) as
    PDF through the configured office converter (Gotenberg or Graph), or
    None when no converter is configured or the conversion fails. The PDF
    is the in-app preview copy; the office file is always stored, so a
    converter outage must not fail the render. That covers every failure,
    not only ConversionError: the Graph converter's upload raises httpx
    errors (HTTPStatusError, TransportError) that convert_for_send does not
    wrap. `filename` picks the converter's options (a spreadsheet renders
    one PDF page per sheet).

    Synchronous and blocking (network I/O, and convert_for_send sleeps two
    seconds before its retry): call it from a worker thread or through
    asyncio.to_thread, never directly inside `async def`."""
    from app.services import office_preview

    try:
        if office_preview.get_converter() is None:
            return None
        return _drop_blank_pages(office_preview.convert_for_send(content, filename))
    except Exception as exc:
        log.warning(
            "AI Estimating PDF conversion of %s failed (%s): %s", filename, type(exc).__name__, exc
        )
        return None


# A converted page narrower or shorter than this (points) holds nothing
# legible: it is the workbook's blank first sheet (kept in the xlsx to match
# the estimator's sent files), which LibreOffice prints as a 48 x 0 pt page.
_BLANK_PAGE_MIN_POINTS = 36.0


def _drop_blank_pages(pdf: bytes) -> bytes:
    """`pdf` without its degenerate pages (see _BLANK_PAGE_MIN_POINTS), or
    unchanged when every page is real, when none is, or when it cannot be
    read: the copy is a preview, so a bad trim must never lose it."""
    try:
        from pypdf import PdfReader, PdfWriter

        reader = PdfReader(io.BytesIO(pdf))
        keep = [
            page for page in reader.pages
            if float(page.mediabox.width) >= _BLANK_PAGE_MIN_POINTS
            and float(page.mediabox.height) >= _BLANK_PAGE_MIN_POINTS
        ]
        if not keep or len(keep) == len(reader.pages):
            return pdf
        writer = PdfWriter()
        for page in keep:
            writer.add_page(page)
        out = io.BytesIO()
        writer.write(out)
        return out.getvalue()
    except Exception as exc:  # noqa: BLE001 - keep the untrimmed copy
        log.warning("AI Estimating PDF copy trim skipped (%s): %s", type(exc).__name__, exc)
        return pdf


def brief_pdf_bytes(docx_bytes: bytes) -> bytes | None:
    """The brief as PDF (office_pdf_bytes with the brief's file name)."""
    return office_pdf_bytes(docx_bytes, "Deep Brief.docx")


# ---------------------------------------------------------------------------
# File names
# ---------------------------------------------------------------------------

_FILENAME_BAD = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def _filename_part(text: str) -> str:
    s = _FILENAME_BAD.sub(" ", _clean(text))
    return re.sub(r"\s+", " ", s).strip().strip(".")


def deliverable_filename(
    kind: str, project_name: str, version: int, *, area: str | None = None
) -> str:
    """The sent-file names: 'BOQ - COUNTS.xlsx', '<UPPER NAME> - ELECTRICAL
    ESTIMATE.xlsx', 'Trench Markup.pdf', 'Sawcut Markup.pdf', 'Marked
    Plans.pdf' or 'Marked Plans - <AREA>.pdf', 'Deep Brief.docx' and (kind
    'brief_pdf') 'Deep Brief.pdf'. Version 2 and later get 'Rev <n> - '."""
    if kind == "boq":
        name = "BOQ - COUNTS.xlsx"
    elif kind == "estimate":
        project = _filename_part(project_name).upper() or "PROJECT"
        name = f"{project[:150]} - ELECTRICAL ESTIMATE.xlsx"
    elif kind == "trench_markup":
        name = "Trench Markup.pdf"
    elif kind == "sawcut_markup":
        name = "Sawcut Markup.pdf"
    elif kind == "marked_plans":
        area_s = _filename_part(area or "").upper()
        name = f"Marked Plans - {area_s[:120]}.pdf" if area_s else "Marked Plans.pdf"
    elif kind == "brief":
        name = "Deep Brief.docx"
    elif kind == "brief_pdf":
        name = "Deep Brief.pdf"
    else:
        raise ValueError(f"Unknown deliverable kind '{_clean(kind)[:40]}'.")
    try:
        v = int(version)
    except (TypeError, ValueError):
        v = 1
    return f"Rev {v} - {name}" if v > 1 else name
