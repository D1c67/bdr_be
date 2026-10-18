"""scripts/plans_reader_golden.py: estimator BOQ parsing, spec flattening,
description normalization and the AI-vs-estimator BOQ scoring."""

from __future__ import annotations

import importlib.util
import io
import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment, Font

_SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "plans_reader_golden.py"
_spec = importlib.util.spec_from_file_location("plans_reader_golden", _SCRIPT)
golden = importlib.util.module_from_spec(_spec)
sys.modules["plans_reader_golden"] = golden
_spec.loader.exec_module(golden)

BACKUP = golden.DEFAULT_BACKUP_DIR


# Helpers ----------------------------------------------------------------------

def _backup_boq(prefix: str) -> bytes:
    if not (BACKUP / "data" / "project_files.json").exists():
        pytest.skip("production backup not on disk")
    try:
        return golden.find_backup_boq(prefix, BACKUP).read_bytes()
    except FileNotFoundError:
        pytest.skip(f"BOQ for {prefix} not in the backup")


def _counts(rows, *, accessories: bool = False) -> dict[str, int]:
    return dict(Counter(r.group_key for r in rows if r.is_accessory == accessories))


def _build_xlsx(sheets: dict[str, list[tuple]]) -> bytes:
    """Workbook in the estimator layout. Each row tuple is (kind, text, qty, unit)
    with kind 'h0' (level 0), 'h1' (group), 'h2' (subhead), 'item', 'acc'."""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"  # the blank first sheet
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        ws.cell(2, 2, "TEST PROJECT\n123 MAIN ST")
        for col, head in enumerate(("SR.\nNO.", "DESCRIPTION", "QUANTITY", "UNIT"), start=1):
            ws.cell(4, col, head).font = Font(bold=True)
        for offset, (kind, text, qty, unit) in enumerate(rows):
            r = 5 + offset
            b = ws.cell(r, 2, text)
            if kind != "acc":
                ws.cell(r, 1, f'=IF(C{r}<>"",1+MAX(A$5:A{r - 1}),"")')
            if kind in ("h0", "h1"):
                b.font = Font(bold=True)
                b.alignment = Alignment(horizontal="center")
            elif kind == "h2":
                b.font = Font(bold=True)
                b.alignment = Alignment(horizontal="left")
            elif kind == "acc":
                b.alignment = Alignment(horizontal="right")
            if qty is not None:
                ws.cell(r, 3, qty)
            if unit is not None:
                ws.cell(r, 4, unit)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _item(desc, qty, unit="EA", *, wastage=False, accessory=False, caveat=None):
    return {
        "description": desc,
        "quantity": qty,
        "unit": unit,
        "sheet_refs": [],
        "caveat": caveat,
        "material_note": None,
        "material_category": "General Material",
        "wastage_applies": wastage,
        "is_accessory": accessory,
    }


def _spec(areas: dict[str, dict[str, dict[str, dict[str, list]]]]) -> dict:
    """{area: {section: {group: {subgroup: [items]}}}} -> BOQ_SPEC."""
    return {
        "title": "TEST",
        "subtitle": None,
        "closing_notes": [],
        "areas": [
            {
                "area": area,
                "notes": [],
                "sections": [
                    {
                        "heading": sec,
                        "note": None,
                        "groups": [
                            {
                                "name": grp,
                                "note": None,
                                "subgroups": [
                                    {"name": sub, "note": None, "items": items}
                                    for sub, items in subs.items()
                                ],
                            }
                            for grp, subs in groups.items()
                        ],
                    }
                    for sec, groups in secs.items()
                ],
            }
            for area, secs in areas.items()
        ],
    }


_EST_SHEET = [
    ("h0", "NEW WORK", None, None),
    ("h1", "WIRING\n(10% wastage is included in the quantities of all conductors/cables, "
           "and conduits in this sheet)", None, None),
    ("h2", "CONDUCTORS (AWG)", None, None),
    ("item", "#12 THHN/THWN Cu", 1100, "FT"),
    ("item", "#10 THHN/THWN Cu", 550, "FT"),
    ("h2", "CONDUITS", None, None),
    ("item", '3/4"C EMT', 1100, "FT"),
    ("acc", '3/4"C EMT Compression Conn Stl', 22, "EA"),
    ("acc", '3/4"C EMT Compression Coupling Stl', 110, "EA"),
    ("h2", "OTHER CONDUIT ACCESSORIES", 100, None),  # stray number, still a heading
    ("acc", "#8 To #10 X 7/8 Plas Anchor (3/16)", 55, "EA"),
    ("h1", "SWITCHGEAR", None, None),
    ("h2", "CIRCUIT BREAKERS", None, None),
    ("item", "20A/1P CIRCUIT BREAKER", 10, "EA"),
    ("h1", "OTHER ITEMS", None, None),
    ("h2", 'RECEPTACLES (RECEPTACLES INCLUDES 4"S DEEP COMB BOX AND COVER PLATE)', None, None),
    ("item", "DUPLEX RECEPTACLE", 20, "EA"),
    ("item", "DUPLEX RECEPTACLE, GFCI", 4, "EA"),
    ("h1", "LIGHTING FIXTURES", None, None),
    ("item", "A, 2X4 LED TROFFER", 12, "EA"),
    ("h1", 'LOW VOLTAGE WORK\n(OUTLETS/DEVICES INCLUDES 4" S BACKBOX, MUD RINGS AND COVER PLATES)',
     None, None),
    ("item", "TELE/DATA OUTLET", 8, "EA"),
]


def _ai_spec() -> dict:
    return _spec({
        "SHEET": {
            "NEW WORK": {
                "WIRING": {
                    "CONDUCTORS (AWG)": [
                        _item("#12 THHN/THWN CU", 1000, "FT", wastage=True),  # 1100, exact
                        _item("#10 THHN THWN Cu", 600, "FT", wastage=True),  # 660, +20%
                    ],
                    "CONDUITS": [
                        _item('3/4" C EMT', 1000, "FT", wastage=True, caveat="(ASSUMED)"),
                        _item('3/4"C EMT Compression Coupling Stl', 110, accessory=True),
                    ],
                },
                "SWITCHGEAR": {
                    "CIRCUIT BREAKERS": [_item("20A, 1P CIRCUIT BREAKER", 10)],
                },
                "OTHER ITEMS": {
                    "RECEPTACLES": [_item("DUPLEX RECEPTACLE", 25)],  # +25%
                    "MISCELLANEOUS": [_item("FLOOR BOX WITH POKE-THRU", 3)],  # not on the BOQ
                },
                "LIGHTING FIXTURES": {"FIXTURES": [_item("A, 2X4 LED TROFFER", 12)]},
            }
        }
    })


# Estimator parsing ------------------------------------------------------------

def test_parse_synthetic_layout():
    rows = golden.parse_estimator_boq(_build_xlsx({"SHEET": _EST_SHEET}))
    assert {r.area for r in rows} == {"SHEET"}  # blank Sheet1 dropped
    assert _counts(rows) == {
        "NEW WORK / WIRING": 3,
        "NEW WORK / SWITCHGEAR": 1,
        "NEW WORK / OTHER ITEMS": 2,
        "NEW WORK / LIGHTING FIXTURES": 1,
        "NEW WORK / LOW VOLTAGE WORK": 1,
    }
    assert _counts(rows, accessories=True) == {"NEW WORK / WIRING": 3}
    by_desc = {r.description: r for r in rows}
    assert by_desc["#12 THHN/THWN Cu"].subgroup == "CONDUCTORS (AWG)"
    assert by_desc["#12 THHN/THWN Cu"].quantity == 1100
    assert by_desc['3/4"C EMT'].subgroup == "CONDUITS"
    anchor = by_desc["#8 To #10 X 7/8 Plas Anchor (3/16)"]
    assert anchor.is_accessory and anchor.subgroup == "OTHER CONDUIT ACCESSORIES"
    assert by_desc["DUPLEX RECEPTACLE"].subgroup.startswith("RECEPTACLES")
    assert not any(r.description == "OTHER CONDUIT ACCESSORIES" for r in rows)


def test_parse_sheet_without_level0_defaults_to_new_work():
    xlsx = _build_xlsx({
        "AREA ONE": [("h0", "DEMOLITION", None, None), ("item", "REMOVAL OF PULLBOX", 2, "EA")],
        "AREA TWO": [
            ("h1", "OTHER ITEMS", None, None),
            ("h2", "MISCELLANEOUS\n(ASSUMED SCOPE, NEED TO CONFIRM)", None, None),
            ("item", "PULLBOX", 3, "EA"),
        ],
    })
    rows = golden.parse_estimator_boq(xlsx)
    assert [(r.area, r.group_key, r.subgroup) for r in rows] == [
        ("AREA ONE", "DEMOLITION / DEMOLITION", ""),
        ("AREA TWO", "NEW WORK / OTHER ITEMS", "MISCELLANEOUS"),
    ]


def test_parse_real_backup_five_star():
    rows = golden.parse_estimator_boq(_backup_boq("dc0cb2f0"))
    assert _counts(rows) == {
        "NEW WORK / WIRING": 14,
        "NEW WORK / SWITCHGEAR": 12,
        "NEW WORK / OTHER ITEMS": 10,
        "NEW WORK / LIGHTING FIXTURES": 13,
        "NEW WORK / LIGHTING CONTROLS": 2,
        "NEW WORK / LOW VOLTAGE WORK": 5,
    }
    assert _counts(rows, accessories=True) == {
        "NEW WORK / WIRING": 12,
        "NEW WORK / LOW VOLTAGE WORK": 6,
    }
    by_desc = {r.description: r for r in rows if not r.is_accessory}
    assert by_desc["#12 THHN/THWN Cu"].subgroup == "CONDUCTORS (AWG)"
    assert by_desc["CAT6 CABLE"].quantity == pytest.approx(13161.06)
    assert by_desc["CAT6 CABLE"].subgroup == "LOW VOLTAGE WIRING"
    assert {r.unit for r in rows} <= set(golden.UNITS)


def test_parse_real_backup_edikted_and_traffic_signal():
    edikted = golden.parse_estimator_boq(_backup_boq("026b2dc5"))
    counts = _counts(edikted)
    assert counts["DEMOLITION / DEMOLITION"] == 4
    assert counts["NEW WORK / WIRING"] == 9
    assert counts["NEW WORK / LIGHTING CONTROLS"] == 11
    assert {r.area for r in edikted} == {"SHEET"}  # Chart1 and Sheet1 dropped

    traffic = golden.parse_estimator_boq(_backup_boq("26ce6d73"))
    assert len({r.area for r in traffic}) == 8
    silverado = [r for r in traffic if r.area == "SILVERADO RANCH"]
    assert _counts(silverado) == {
        "DEMOLITION / DEMOLITION": 3,
        "NEW WORK / WIRING": 20,
        "NEW WORK / OTHER ITEMS": 17,
    }


# Spec flattening --------------------------------------------------------------

def test_rows_from_spec_applies_wastage_and_caveats():
    spec = _spec({
        "SHEET": {
            "NEW WORK": {
                "WIRING": {
                    "CONDUCTORS (AWG)": [
                        _item("#12 THHN/THWN Cu", 100, "FT", wastage=True),
                        _item("#10 THHN/THWN Cu", 100, "FT", wastage=False),
                    ],
                },
                "OTHER ITEMS": {"MISCELLANEOUS": [
                    _item("4\"S JUNCTION BOX", 10, wastage=True, caveat="(ASSUMED)"),
                ]},
            },
            "DEMOLITION": {"REMOVALS": {"-": [_item("REMOVAL OF PANEL", 1)]}},
        }
    })
    rows = golden.rows_from_spec(spec)
    assert [r.quantity for r in rows] == [pytest.approx(110), 100, 10, 1]
    assert rows[2].description == '4"S JUNCTION BOX (ASSUMED)'
    assert rows[3].group_key == "DEMOLITION / DEMOLITION"
    # The full main agent output is accepted too.
    assert len(golden.rows_from_spec({"boq": spec, "brief": {}})) == 4


# Normalization ----------------------------------------------------------------

@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ('3/4" C EMT', '3/4"C EMT'),
        ('3/4"C EMT', '3/4"C EMT'),
        ("3/4\u201d c emt", '3/4"C EMT'),
        ('\u00be" C EMT', '3/4"C EMT'),
        ('1 1/4" C EMT', '1-1/4"C EMT'),
        ('1\u00bc"C EMT', '1-1/4"C EMT'),
        ("3/4 IN C EMT", '3/4"C EMT'),
        ("DUPLEX RECEPTACLE (ASSUMED).", "DUPLEX RECEPTACLE"),
        ("NOEN SWITCH (PLEASE CONFIRM)", "NOEN SWITCH"),
        ("HOUSE PANEL 100A (RATING DETAILS NOT GIVEN, NEED TO CONFIRM SCOPE OF WORK)",
         "HOUSE PANEL 100A"),
        ("#8 To #10 X 7/8 Plas Anchor (3/16)", "#8 TO #10 X 7/8 PLAS ANCHOR (3/16)"),
        ("  e3 ,  ceiling   mount\nlight ;", "E3, CEILING MOUNT LIGHT"),
        ("SWITCHBOARD \u2018MSA\u2019", "SWITCHBOARD 'MSA'"),
        (None, ""),
    ],
)
def test_normalize_description(raw, expected):
    assert golden.normalize_description(raw) == expected


def test_key_tokens_tell_sizes_apart():
    def row(d, u="FT"):
        return golden.BoqRow("S", "NEW WORK", "WIRING", "", d, 1, u)

    assert golden.similarity(row("#12 THHN/THWN Cu"), row("#12 AWG THHN COPPER")) >= 0.72
    assert golden.similarity(row("#12 THHN/THWN Cu"), row("#10 THHN/THWN Cu")) < 0.72
    assert golden.similarity(row('3/4"C EMT'), row('3/4"C PVC')) < 0.72
    assert golden.similarity(row('3/4"C EMT'), row('1"C EMT')) < 0.72
    assert golden.similarity(row("20A/1P CIRCUIT BREAKER", "EA"),
                             row("20A 1-POLE BREAKER", "EA")) >= 0.72
    assert golden.similarity(row("20A/1P CIRCUIT BREAKER", "EA"),
                             row("20A/3P CIRCUIT BREAKER", "EA")) < 0.72
    assert golden.similarity(row("DUPLEX RECEPTACLE", "EA"),
                             row("DUPLEX RECEPTACLE, GFCI", "EA")) < 0.72
    assert golden.similarity(row("E3, CEILING MOUNT LIGHT", "EA"),
                             row("E4, CEILING MOUNT LIGHT", "EA")) < 0.72


# Matching ---------------------------------------------------------------------

def test_synthetic_spec_vs_synthetic_xlsx():
    est = golden.parse_estimator_boq(_build_xlsx({"SHEET": _EST_SHEET}))
    ai = golden.rows_from_spec(_ai_spec())
    rep = golden.match_rows(ai, est)

    assert rep.ignored_accessories == {"ai": 1, "est": 3}
    assert rep.ai_total == 7 and rep.est_total == 8
    assert len(rep.pairs) == 6
    assert rep.precision == pytest.approx(6 / 7, abs=1e-4)
    assert rep.recall == pytest.approx(6 / 8, abs=1e-4)
    # Errors over matched rows: 0, 20, 0, 0, 25, 0 percent.
    assert rep.mape == pytest.approx(7.5, abs=0.01)
    assert rep.mape_rows == 6
    assert rep.exact_matches == 4
    assert rep.unit_mismatches == 0
    assert rep.group_mismatches == 0
    assert [r.description for r in rep.unmatched_ai] == ["FLOOR BOX WITH POKE-THRU"]
    assert sorted(r.description for r in rep.unmatched_est) == [
        "DUPLEX RECEPTACLE, GFCI", "TELE/DATA OUTLET",
    ]
    pg = rep.per_group
    assert pg["NEW WORK / WIRING"]["recall"] == 1.0
    assert pg["NEW WORK / OTHER ITEMS"]["recall"] == 0.5
    assert pg["NEW WORK / OTHER ITEMS"]["precision"] == 0.5
    assert pg["NEW WORK / LOW VOLTAGE WORK"]["recall"] == 0.0
    wire10 = next(p for p in rep.pairs if p.est.description == "#10 THHN/THWN Cu")
    assert wire10.ai.quantity == pytest.approx(660)
    assert wire10.abs_error == pytest.approx(110)
    assert wire10.pct_error == pytest.approx(20.0)
    assert not wire10.exact

    # Scoring the accessories as well adds the coupling pair.
    rep_all = golden.match_rows(ai, est, ignore_accessories=False)
    assert rep_all.ai_total == 8 and rep_all.est_total == 11
    assert len(rep_all.pairs) == 7


def test_unit_mismatch_is_flagged_and_left_out_of_mape():
    est = [
        golden.BoqRow("S", "NEW WORK", "LIGHTING FIXTURES", "", "A, 2X4 LED TROFFER", 12, "EA"),
        golden.BoqRow("S", "NEW WORK", "LIGHTING FIXTURES", "", "TRACK ROD", 300, "FT"),
    ]
    ai = [
        golden.BoqRow("S", "NEW WORK", "LIGHTING FIXTURES", "", "A, 2X4 LED TROFFER", 48, "FT"),
        golden.BoqRow("S", "NEW WORK", "LIGHTING FIXTURES", "", "TRACK ROD", 330, "LF"),
    ]
    rep = golden.match_rows(ai, est)
    assert len(rep.pairs) == 2 and rep.unit_mismatches == 1
    troffer = next(p for p in rep.pairs if p.est.unit == "EA")
    assert troffer.unit_mismatch and troffer.pct_error is None and not troffer.exact
    rod = next(p for p in rep.pairs if p.est.unit == "FT")
    assert not rod.unit_mismatch  # LF normalizes to FT
    assert rod.pct_error == pytest.approx(10.0)
    assert rep.mape == pytest.approx(10.0) and rep.mape_rows == 1
    assert "UNIT EA vs FT" in golden.report_markdown(rep)


def test_cross_group_match_and_areas():
    est = [
        golden.BoqRow("NORTH LOT", "NEW WORK", "OTHER ITEMS", "", "PULLBOX", 4, "EA"),
        golden.BoqRow("SOUTH LOT", "NEW WORK", "OTHER ITEMS", "", "PULLBOX", 9, "EA"),
        golden.BoqRow("SOUTH LOT", "NEW WORK", "SWITCHGEAR", "", "200A SERVICE PEDESTAL", 1, "EA"),
    ]
    ai = [
        golden.BoqRow("SOUTH LOT", "NEW WORK", "OTHER ITEMS", "", "PULL BOX", 9, "EA"),
        golden.BoqRow("NORTH LOT AREA", "NEW WORK", "OTHER ITEMS", "", "PULLBOX", 4, "EA"),
        golden.BoqRow("SOUTH LOT", "NEW WORK", "OTHER ITEMS", "", "200A SERVICE PEDESTAL", 1, "EA"),
    ]
    rep = golden.match_rows(ai, est)
    assert rep.precision == 1.0 and rep.recall == 1.0 and rep.exact_matches == 3
    assert {(p.ai.area, p.est.area) for p in rep.pairs if "PULL" in p.est.description} == {
        ("SOUTH LOT", "SOUTH LOT"), ("NORTH LOT AREA", "NORTH LOT"),
    }
    assert rep.group_mismatches == 1


# Reports and CLI --------------------------------------------------------------

def test_reports_and_cli(tmp_path, capsys):
    xlsx = tmp_path / "BOQ - COUNTS.xlsx"
    xlsx.write_bytes(_build_xlsx({"SHEET": _EST_SHEET}))
    spec = tmp_path / "spec.json"
    spec.write_text(json.dumps(_ai_spec()))
    out = tmp_path / "report.json"

    assert golden.main(["--spec", str(spec), "--estimator", str(xlsx), "--json", str(out)]) == 0
    md = capsys.readouterr().out
    assert "| Precision | 85.7% |" in md
    assert "| Recall | 75.0% |" in md
    assert "TELE/DATA OUTLET" in md and "FLOOR BOX WITH POKE-THRU" in md
    assert "\u2014" not in md

    data = json.loads(out.read_text())
    assert data["totals"]["matched"] == 6
    assert data["totals"]["mape"] == pytest.approx(7.5, abs=0.01)
    assert len(data["pairs"]) == 6 and len(data["unmatched_est"]) == 2
    assert data["per_group"]["NEW WORK / LOW VOLTAGE WORK"]["recall"] == 0.0

    assert golden.main(["--estimator", str(xlsx)]) == 0
    assert "NEW WORK / WIRING | 3 | 3" in capsys.readouterr().out


def test_normalize_description_keeps_only_the_first_line():
    """A rendered AI row puts its material note on a second line; the note
    must not drag the text similarity below the match threshold."""
    from scripts.plans_reader_golden import normalize_description

    two_line = "#3/0 THHN/THWN Cu (ASSUMED)\n(Feeder 1: 4 x 300 FT; feeder 3: 4 x 10 FT.)"
    assert normalize_description(two_line) == normalize_description("#3/0 THHN/THWN Cu")
    assert normalize_description("PANEL A, 277/480V") == "PANEL A, 277/480V"
