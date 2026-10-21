"""Auto send estimate parser (docs/AUTO_SEND.md section 4).

Workbooks are built here with openpyxl in the real estimator layout (title
B2, headers on row 3, line rows, TOTALS two rows below the last line, BID
SUMMARY three rows below that, BASE BID PRICE last). openpyxl does not save
cached formula values, so the priced fixtures write LITERAL numbers where
Excel would have saved them; the formula-only fixture writes the formulas
to exercise no_cached_values. No real customer workbook is used.
"""

from __future__ import annotations

import datetime
import io
import random
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from app.services import auto_send_estimate as ase
from app.services import plans_reader_render as prr

D = Decimal
# Em and en dash, never allowed in owner-facing text.
DASHES = {chr(0x2014), chr(0x2013)}

HEADERS = (
    "SR. NO.",
    "DESCRIPTION",
    "MATERIAL COST",
    "LABOR COST",
    "MATERIAL TAX",
    "LABOR TAX",
    "TOTAL COST",
    "OHP (LABOR)",
    "OHP (MATERIAL)",
    "TOTAL PRICE",
)


# ── fixture builder ────────────────────────────────────────────────────────


def recap_book(
    lines=(("WIRING", 10000.0, 20000.0), ("OTHER ITEMS", 5000.0, 4000.0)),
    *,
    sheet="Bid Recap & Summary",
    areas=("Detailed Estimate",),
    header_row=3,
    first_col=2,
    mat_tax=0.0,
    lab_tax=0.0,
    ohp_labor=0.25,
    ohp_material=0.25,
    bid_sec_rate=0.0,
    formulas=False,
    overrides=None,
    drop_labels=(),
    chart=True,
) -> bytes:
    """A recap workbook. `overrides` maps a label ('TOTALS:<key>' for a
    TOTALS cost cell, or a summary label) to the value to write instead."""
    overrides = overrides or {}
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    if chart:
        # The real workbooks open with a "Chart1" chartsheet (no cells).
        wb.active["A1"] = 1
        bar = BarChart()
        bar.add_data(Reference(wb.active, min_col=1, min_row=1, max_row=1))
        wb.create_chartsheet("Chart1").add_chart(bar)
    for name in areas:
        ws_area = wb.create_sheet(name)
        ws_area["P17"] = 1.0
    ws = wb.create_sheet(sheet)
    c0 = first_col  # SR. NO. column; DESCRIPTION is c0 + 1, costs from c0 + 2
    col = {
        k: c0 + 2 + i for i, k in enumerate(("mc", "lc", "mt", "lt", "tc", "ohpl", "ohpm", "tp"))
    }
    letter = {k: openpyxl.utils.get_column_letter(v) for k, v in col.items()}
    ws.cell(row=header_row - 1, column=c0, value="BID RECAP")
    for i, h in enumerate(HEADERS):
        ws.cell(row=header_row, column=c0 + i, value=h)

    first = header_row + 1
    totals = {k: 0.0 for k in col}
    for i, (desc, mc, lc) in enumerate(lines):
        r = first + i
        mt, lt = mc * mat_tax, lc * lab_tax
        tc = mc + lc + mt + lt
        ohpl, ohpm = lc * ohp_labor, mc * ohp_material
        vals = {"mc": mc, "lc": lc, "mt": mt, "lt": lt, "tc": tc, "ohpl": ohpl, "ohpm": ohpm}
        vals["tp"] = tc + ohpl + ohpm
        ws.cell(row=r, column=c0, value=i + 1)
        ws.cell(row=r, column=c0 + 1, value=desc)
        for k, v in vals.items():
            ws.cell(row=r, column=col[k], value=f"=0+{v}" if formulas else v)
            totals[k] += v
    last = first + len(lines) - 1
    tot_row = last + 2
    ws.cell(row=tot_row, column=c0 + 1, value="TOTALS")
    for k in col:
        v = overrides.get(f"TOTALS:{k}", totals[k])
        if formulas:
            v = f"=SUM({letter[k]}{first}:{letter[k]}{last})"
        ws.cell(row=tot_row, column=col[k], value=v)

    tp = totals["tp"]
    bid_sec = tp * bid_sec_rate
    summary = [
        ("TOTAL MATERIAL COST", None, totals["mc"]),
        ("MATERIAL SALES TAX", mat_tax, totals["mt"]),
        ("TOTAL LABOR COST", None, totals["lc"]),
        ("LABOR TAX", lab_tax, totals["lt"]),
        ("TOTAL COST", None, totals["tc"]),
        ("OVERHEAD & PROFIT (LABOR) @", ohp_labor, totals["ohpl"]),
        ("OVERHEAD & PROFIT (MATERIAL) @", ohp_material, totals["ohpm"]),
        ("TOTAL COST WITH OVERHEADS + PROFIT", None, tp),
        ("BID SECURITY", bid_sec_rate, bid_sec),
    ]
    rate_col, value_col = c0 + 2, c0 + 3
    srow = tot_row + 3
    ws.cell(row=srow, column=c0, value="BID SUMMARY")
    r = srow + 1
    for label, rate, value in summary:
        if label not in drop_labels:
            ws.cell(row=r, column=c0 + 1, value=label)
            if rate is not None:
                ws.cell(row=r, column=rate_col, value=rate)
            v = overrides.get(label, value)
            ws.cell(row=r, column=value_col, value=f"=0+{v}" if formulas else v)
        r += 1
    base_row = r + 1
    if "BASE BID PRICE" not in drop_labels:
        ws.cell(row=base_row, column=c0, value=5)
        ws.cell(row=base_row, column=c0 + 1, value="BASE BID PRICE")
        v = overrides.get("BASE BID PRICE", tp + bid_sec)
        ws.cell(row=base_row, column=value_col, value="=SUM(E1:E2)" if formulas else v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── parse_recap ────────────────────────────────────────────────────────────


def test_single_area_reads_materials_and_labor():
    out = ase.parse_recap(recap_book())
    # materials = mc + mt + ohpm = 15000 + 0 + 3750; labor = 24000 + 0 + 6000
    assert out["sheet"] == "Bid Recap & Summary"
    assert out["materials"] == D("18750.00")
    assert out["labor"] == D("30000.00")
    assert out["base"] == D("48750.00")
    assert out["materials"] + out["labor"] == out["base"]
    comp = out["components"]
    assert set(comp) == {"mc", "lc", "mt", "lt", "tc", "ohpl", "ohpm", "tp", "bid_sec", "base"}
    assert comp["mc"] == D("15000.0") and comp["lc"] == D("24000.0")
    assert comp["tp"] == D("48750.0") and comp["bid_sec"] == 0
    assert out["warnings"] == []


def test_real_layout_values_round_like_the_survey():
    # Anita Gelato in the prod survey: materials 42,020.66, labor 60,848.92.
    mc, lc = 42020.66 / 1.25, 60848.92 / 1.25
    out = ase.parse_recap(recap_book(lines=(("WIRING", mc, lc),)))
    assert out["materials"] == D("42020.66")
    assert out["labor"] == D("60848.92")
    assert out["base"] == D("102869.58")


def test_multi_area_and_shifted_layout_found_by_label():
    lines = [(f"AREA {i} WIRING", 1000.0 * i, 500.0 * i) for i in range(1, 9)]
    data = recap_book(
        lines=lines,
        sheet="BID RECAP",  # not the exact name: first sheet containing "recap"
        areas=("BLD A", "BLD B", "BLD C"),
        header_row=6,
        first_col=3,
    )
    out = ase.parse_recap(data)
    mc, lc = sum(x[1] for x in lines), sum(x[2] for x in lines)
    assert out["sheet"] == "BID RECAP"
    assert out["materials"] == ase.q(D(str(mc * 1.25)))
    assert out["labor"] == ase.q(D(str(lc * 1.25)))
    assert out["materials"] + out["labor"] == out["base"]


def test_exact_recap_name_wins_over_other_recap_sheets():
    data = recap_book()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb.create_sheet("Old recap", 0)  # an empty decoy before the real sheet
    buf = io.BytesIO()
    wb.save(buf)
    assert ase.parse_recap(buf.getvalue())["sheet"] == "Bid Recap & Summary"


def test_materials_zero_is_valid_labor_only():
    out = ase.parse_recap(recap_book(lines=(("LABOR ONLY", 0.0, 5108.4),)))
    assert out["materials"] == D("0.00")
    assert out["labor"] == D("6385.50")
    assert out["base"] == D("6385.50")
    assert any("no material cost" in w for w in out["warnings"])


def test_bid_security_spread_so_parts_equal_base():
    # tp = 1234.57*1.25 + 2345.67*1.25 = 4475.30; 1.5% bid security = 67.1295
    data = recap_book(lines=(("WIRING", 1234.57, 2345.67),), bid_sec_rate=0.015)
    out = ase.parse_recap(data)
    comp = out["components"]
    assert comp["bid_sec"] > 0
    assert out["base"] == ase.q(comp["base"])
    assert out["materials"] + out["labor"] == out["base"]
    mat_raw = comp["mc"] + comp["mt"] + comp["ohpm"]
    expect_m = ase.q(mat_raw + comp["bid_sec"] * mat_raw / comp["tp"])
    assert out["materials"] == expect_m
    assert out["labor"] == out["base"] - expect_m
    assert any("Bid security" in w for w in out["warnings"])


def test_ref_error_is_formula_error():
    data = recap_book(overrides={"TOTALS:lc": "#REF!"})
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(data)
    assert e.value.code == "formula_error"
    assert e.value.message == ase.REASON_MESSAGES["formula_error"]


def test_text_in_price_cell_is_formula_error():
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(recap_book(overrides={"BASE BID PRICE": "TBD"}))
    assert e.value.code == "formula_error"


def test_formula_only_workbook_is_no_cached_values():
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(recap_book(formulas=True))
    assert e.value.code == "no_cached_values"


def test_unpriced_ai_estimating_workbook():
    item = {
        "description": "DUPLEX RECEPTACLE",
        "quantity": 10,
        "unit": "EA",
        "sheet_refs": ["E-101"],
        "caveat": None,
        "material_note": None,
        "material_category": "General Material",
        "wastage_applies": False,
        "is_accessory": False,
    }
    boq = {
        "title": "TEST PROJECT",
        "subtitle": "LAS VEGAS, NV",
        "areas": [
            {
                "area": "MAIN",
                "notes": [],
                "sections": [
                    {
                        "heading": "NEW WORK",
                        "note": None,
                        "groups": [
                            {
                                "name": "OTHER ITEMS",
                                "note": None,
                                "subgroups": [{"name": "", "note": None, "items": [item]}],
                            }
                        ],
                    }
                ],
            }
        ],
        "closing_notes": [],
    }
    est = {
        "project_name": "TEST PROJECT",
        "address": "LAS VEGAS, NV",
        "include_standard_allowances": True,
        "extra_allowances": [],
        "general_notes": [],
        "area_notes": [],
    }
    data = prr.render_estimate_xlsx(boq, est, date=datetime.date(2026, 10, 19))
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(data)
    assert e.value.code == "unpriced_or_zero"


def test_zero_placeholder_recap_is_unpriced():
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(recap_book(lines=(("WIRING", 0.0, 0.0),)))
    assert e.value.code == "unpriced_or_zero"


def test_no_recap_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Detailed Estimate")["A1"] = "ESTIMATE"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(buf.getvalue())
    assert e.value.code == "no_recap_sheet"


def test_unreadable_bytes():
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(b"%PDF-1.7 not a workbook")
    assert e.value.code == "unreadable"


def test_missing_header_totals_and_base():
    wb = openpyxl.load_workbook(io.BytesIO(recap_book()))
    wb["Bid Recap & Summary"]["F3"] = "SOMETHING ELSE"  # MATERIAL TAX header gone
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(buf.getvalue())
    assert e.value.code == "no_header"

    wb = openpyxl.load_workbook(io.BytesIO(recap_book()))
    ws = wb["Bid Recap & Summary"]
    for row in ws.iter_rows():
        for c in row:
            if c.value == "TOTALS":
                c.value = "SUM"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(buf.getvalue())
    assert e.value.code == "no_totals_row"

    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(recap_book(drop_labels=("BASE BID PRICE",)))
    assert e.value.code == "no_base_bid"


def test_crosscheck_summary_material_cost_overwritten():
    # Estimator typed over TOTAL MATERIAL COST: it no longer equals TOTALS mc.
    data = recap_book(overrides={"TOTAL MATERIAL COST": 14000.0})
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(data)
    assert e.value.code == "crosscheck_failed"
    assert "TOTAL MATERIAL COST" in e.value.message
    assert DASHES.isdisjoint(e.value.message)


def test_crosscheck_total_cost_and_base_mismatch():
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(recap_book(overrides={"TOTALS:tc": 1.0}))
    assert e.value.code == "crosscheck_failed"
    with pytest.raises(ase.EstimateError) as e:
        ase.parse_recap(recap_book(overrides={"BASE BID PRICE": 60000.0}))
    assert e.value.code == "crosscheck_failed"


def test_crosscheck_tolerates_float_noise():
    out = ase.parse_recap(recap_book(overrides={"BASE BID PRICE": 48750.0 + 0.04}))
    assert out["base"] == D("48750.04")
    assert out["materials"] + out["labor"] == out["base"]


def test_summary_rows_optional_except_base():
    out = ase.parse_recap(
        recap_book(drop_labels=("TOTAL MATERIAL COST", "TOTAL LABOR COST", "BID SECURITY"))
    )
    assert out["base"] == D("48750.00")


def test_unusual_rate_is_warning_not_failure():
    out = ase.parse_recap(recap_book(ohp_labor=1.5))
    assert out["materials"] + out["labor"] == out["base"]
    assert any("labor overhead and profit" in w for w in out["warnings"])


def test_reason_messages_plain_text():
    for code, msg in ase.REASON_MESSAGES.items():
        assert msg and DASHES.isdisjoint(msg), code


# ── apply_markup ───────────────────────────────────────────────────────────


def test_apply_markup_ten_percent_rounding():
    m, lab, total = ase.apply_markup(D("42020.66"), D("60848.92"), D("10"))
    assert (m, lab, total) == (D("46222.73"), D("66933.81"), D("113156.54"))


def test_apply_markup_half_up_and_zero_pct():
    assert ase.apply_markup(D("0.05"), D("0.15"), D("10")) == (D("0.06"), D("0.17"), D("0.23"))
    assert ase.apply_markup(D("100.00"), D("0"), D("0")) == (D("100.00"), D("0.00"), D("100.00"))
    assert ase.apply_markup(D("1000"), D("2000"), D("12.5")) == (
        D("1125.00"),
        D("2250.00"),
        D("3375.00"),
    )


def test_apply_markup_total_is_always_sum_of_parts():
    rng = random.Random(159)
    for _ in range(2000):
        m = D(rng.randint(0, 10_000_000)) / 100
        lab = D(rng.randint(0, 10_000_000)) / 100
        pct = D(rng.randint(0, 10000)) / 100
        mm, ll, tt = ase.apply_markup(m, lab, pct)
        assert tt == mm + ll
        assert mm == mm.quantize(D("0.01")) and ll == ll.quantize(D("0.01"))


# ── fake Supabase ──────────────────────────────────────────────────────────


class _Q:
    def __init__(self, sb, table):
        self.sb, self.table = sb, table
        self.filters = []
        self.op, self.payload, self.lim = "select", None, None

    def select(self, *a, **k):
        return self

    def eq(self, col, val):
        self.filters.append(lambda r: r.get(col) == val)
        return self

    @property
    def not_(self):
        q = self

        class _N:
            def is_(self, col, val):
                assert val == "null"
                q.filters.append(lambda r: r.get(col) is not None)
                return q

        return _N()

    def limit(self, n):
        self.lim = n
        return self

    def upsert(self, payload, on_conflict=None):
        self.op, self.payload, self.on_conflict = "upsert", payload, on_conflict
        return self

    def execute(self):
        if self.table in self.sb.fail_tables:
            raise RuntimeError(f"{self.table} down")
        rows = self.sb.tables.setdefault(self.table, [])
        self.sb.calls.append((self.table, self.op))
        if self.op == "upsert":
            key = self.on_conflict
            rows[:] = [r for r in rows if r.get(key) != self.payload[key]]
            rows.append(dict(self.payload))
            return SimpleNamespace(data=[dict(self.payload)])
        hits = [dict(r) for r in rows if all(f(r) for f in self.filters)]
        if self.lim is not None:
            hits = hits[: self.lim]
        return SimpleNamespace(data=hits)


class FakeSB:
    def __init__(self, tables=None):
        self.tables = {k: [dict(r) for r in v] for k, v in (tables or {}).items()}
        self.calls = []
        self.fail_tables = set()

    def table(self, name):
        return _Q(self, name)


def _file(fid, filename, rnd, *, deliverable=True, created="2026-10-01T10:00:00+00:00", **kw):
    row = {
        "id": fid,
        "project_id": "p1",
        "category": "estimate",
        "filename": filename,
        "storage_path": f"p1/{fid}/{filename}",
        "size_bytes": 1000,
        "estimator_deliverable": deliverable,
        "submission_round": rnd,
        "uploaded_by": "est-1",
        "sha256": None,
        "created_at": created,
    }
    row.update(kw)
    return row


# ── pick_estimate_file ─────────────────────────────────────────────────────


def test_pick_highest_sealed_round_ignoring_internal_and_drafts():
    sb = FakeSB(
        {
            "project_files": [
                _file("r1", "JOB - ELECTRICAL ESTIMATE.xlsx", 1),
                _file("r2", "Rev - JOB - ELECTRICAL ESTIMATE.XLSX", 2),
                # newest rows are not the estimator's sealed workbook:
                _file("pdf", "RFI 8-6-26.pdf", None, deliverable=False, created="2026-10-09"),
                _file("lv", "7149 Low Voltage.xlsx", None, deliverable=False, created="2026-10-09"),
                _file("draft", "Rev2 - JOB.xlsx", None, created="2026-10-10"),
                _file("other", "x.xlsx", 3, project_id="p2"),
                _file("boq", "boq.xlsx", 3, category="boq"),
            ]
        }
    )
    f = ase.pick_estimate_file(sb, "p1")
    assert f["id"] == "r2"
    assert f["submission_round"] == 2
    assert f["newer_round_without_workbook"] is None
    assert {"id", "filename", "storage_path", "size_bytes", "uploaded_by", "sha256"} <= set(f)


def test_pick_blocks_when_the_latest_round_estimate_is_not_a_workbook():
    """Finding 19 (owner decision: latest SEALED round only): a revised
    estimate sent as a PDF supersedes the older workbook, so the older price is
    never used as a fallback."""
    sb = FakeSB(
        {
            "project_files": [
                _file("r1", "ESTIMATE.xlsm", 1),
                _file("pdf", "Estimate rev.pdf", 2),
            ]
        }
    )
    with pytest.raises(ase.EstimateError) as exc:
        ase.pick_estimate_file(sb, "p1")
    assert exc.value.code == "newer_estimate_not_workbook"
    assert "round 2" in exc.value.message and "round 1" in exc.value.message
    assert "\u2014" not in exc.value.message
    assert "newer_estimate_not_workbook" in ase.REASON_MESSAGES


def test_pick_uses_the_latest_round_workbook_when_it_has_one():
    sb = FakeSB(
        {
            "project_files": [
                _file("r1", "ESTIMATE.xlsm", 1),
                _file("r2", "ESTIMATE rev.xlsx", 2),
                _file("png", "screenshot.png", 2),
            ]
        }
    )
    f = ase.pick_estimate_file(sb, "p1")
    assert f["id"] == "r2"
    assert f["newer_round_without_workbook"] is None


@pytest.mark.parametrize(
    ("files", "code"),
    [
        ([], "estimate_not_received"),
        ([_file("pdf", "RFI.pdf", None, deliverable=False)], "estimate_not_received"),
        ([_file("draft", "ESTIMATE.xlsx", None)], "estimate_not_received"),
        ([_file("pdf", "ESTIMATE.pdf", 1)], "no_estimate_workbook"),
        (
            [_file("a", "A.xlsx", 2), _file("b", "B.xlsx", 2), _file("c", "C.xlsx", 1)],
            "ambiguous_workbooks",
        ),
    ],
)
def test_pick_failures(files, code):
    with pytest.raises(ase.EstimateError) as e:
        ase.pick_estimate_file(FakeSB({"project_files": files}), "p1")
    assert e.value.code == code


# ── load_estimate ──────────────────────────────────────────────────────────


@pytest.fixture
def env(monkeypatch):
    state = SimpleNamespace(downloads=[], blob=recap_book(), fail_download=False)

    def download(path):
        state.downloads.append(path)
        if state.fail_download:
            raise RuntimeError("storage down")
        return state.blob

    monkeypatch.setattr(ase.storage, "download_file", download)
    monkeypatch.setattr(ase, "get_settings", lambda: SimpleNamespace(boq_max_bytes=10_000_000))
    state.sb = FakeSB(
        {
            "project_files": [_file("r1", "JOB - ELECTRICAL ESTIMATE.xlsx", 1)],
            "profiles": [{"id": "est-1", "full_name": "Pat Estimator"}],
        }
    )
    return state


def test_load_estimate_parses_then_uses_cache(env):
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is True and res["code"] is None
    assert res["file_id"] == "r1" and res["round"] == 1
    assert res["uploaded_by"] == "est-1" and res["estimator_name"] == "Pat Estimator"
    assert (res["materials"], res["labor"], res["base"]) == ("18750.00", "30000.00", "48750.00")
    assert res["components"]["mc"] == "15000.00"
    assert isinstance(res["components"]["bid_sec"], str)
    cache = env.sb.tables["auto_send_estimates"]
    assert len(cache) == 1 and cache[0]["estimate_file_id"] == "r1" and cache[0]["ok"] is True
    assert cache[0]["project_id"] == "p1" and cache[0]["sha256"]
    assert len(env.downloads) == 1

    again = ase.load_estimate("p1", sb=env.sb)
    assert len(env.downloads) == 1  # served from the cache
    assert again["materials"] == res["materials"] and again["base"] == res["base"]


def test_load_estimate_reparses_when_sha_changes(env):
    env.sb.tables["project_files"][0]["sha256"] = "aaa"
    ase.load_estimate("p1", sb=env.sb)
    env.sb.tables["project_files"][0]["sha256"] = "bbb"
    env.blob = recap_book(lines=(("WIRING", 100.0, 100.0),))
    res = ase.load_estimate("p1", sb=env.sb)
    assert len(env.downloads) == 2
    assert res["base"] == "250.00"
    assert env.sb.tables["auto_send_estimates"][0]["sha256"] == "bbb"


def test_load_estimate_ignores_cache_from_older_parser(env):
    env.sb.tables["auto_send_estimates"] = [
        {
            "estimate_file_id": "r1",
            "project_id": "p1",
            "sha256": None,
            "ok": True,
            "code": None,
            "message": None,
            "result": {"parser_version": 0, "materials": "1.00", "labor": "1.00", "base": "2.00"},
        }
    ]
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["base"] == "48750.00" and len(env.downloads) == 1


def test_load_estimate_caches_parse_failures(env):
    env.blob = recap_book(formulas=True)
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is False and res["code"] == "no_cached_values"
    assert res["message"] == ase.REASON_MESSAGES["no_cached_values"]
    assert res["file_id"] == "r1" and res["estimator_name"] == "Pat Estimator"
    again = ase.load_estimate("p1", sb=env.sb)
    assert again["code"] == "no_cached_values" and len(env.downloads) == 1


def test_load_estimate_download_failure_not_cached(env):
    env.fail_download = True
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is False and res["code"] == "download_failed"
    assert "auto_send_estimates" not in env.sb.tables or not env.sb.tables["auto_send_estimates"]
    env.fail_download = False
    assert ase.load_estimate("p1", sb=env.sb)["ok"] is True


def test_load_estimate_size_guard_before_download(env):
    env.sb.tables["project_files"][0]["size_bytes"] = 20_000_000
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["code"] == "too_large" and env.downloads == []


def test_load_estimate_no_file_and_db_errors_never_raise(env):
    env.sb.tables["project_files"] = []
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is False and res["code"] == "estimate_not_received"
    assert res["file_id"] is None and res["components"] == {} and res["warnings"] == []

    env.sb.fail_tables.add("project_files")
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is False and res["code"] == "lookup_failed"


def test_load_estimate_cache_write_failure_still_returns_result(env):
    env.sb.fail_tables.add("auto_send_estimates")
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is True and res["base"] == "48750.00"


def test_load_estimate_refuses_a_stale_round_workbook(env):
    env.sb.tables["project_files"].append(_file("png", "cost estimate.png", 2))
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["ok"] is False
    assert res["code"] == "newer_estimate_not_workbook"
    assert res["materials"] is None and res["labor"] is None


def test_load_estimate_carries_the_impersonation_stamp(env):
    env.sb.tables["project_files"][0]["impersonated_by"] = "u-dev"
    res = ase.load_estimate("p1", sb=env.sb)
    assert res["impersonated_by"] == "u-dev"
    assert res["uploaded_at"] == env.sb.tables["project_files"][0].get("created_at")


# ── migration 0159 (SQL text) ──────────────────────────────────────────────


def _migration() -> str:
    from pathlib import Path

    migrations = Path(__file__).resolve().parents[1] / "supabase/migrations"
    (path,) = list(migrations.glob("0159_*.sql"))
    return path.read_text(encoding="utf-8")


AUTO_SEND_TABLES = (
    "auto_send_settings",
    "auto_send_runs",
    "auto_send_gcs",
    "auto_send_email_log",
    "auto_send_estimates",
)


def test_migration_tables_rls_and_house_style():
    sql = _migration()
    assert sql.startswith("-- 0159: Auto send")
    assert sql.rstrip().endswith("notify pgrst, 'reload schema';")
    assert DASHES.isdisjoint(sql)
    for t in AUTO_SEND_TABLES:
        assert f"create table if not exists {t} (" in sql
        assert f"alter table {t} enable row level security;" in sql
        assert f"alter table {t} force row level security;" in sql
    assert "create policy" not in sql
    for t in ("auto_send_settings", "auto_send_runs", "auto_send_gcs"):
        assert f"before update on {t}" in sql
    assert "insert into auto_send_settings (id) values (1) on conflict (id) do nothing;" in sql
    assert "unique (project_id, bid_date)" in sql
    assert "unique (run_id, gc_id)" in sql
    assert "unique (send_date, kind, user_id)" in sql
    assert "estimate_file_id  uuid primary key references project_files(id)" in sql
    assert "check (origin in ('send_out', 'external_submission', 'auto_send'))" in sql
    assert "drop constraint if exists proposal_sends_origin_check" in sql
    assert "references auto_send_runs(id) on delete set null" in sql


def test_migration_acceptance_function_rechecks_under_locks():
    sql = _migration()
    fn = sql[sql.index("create or replace function apply_auto_send_acceptance") :]
    assert "security definer" in fn and "set search_path = public, pg_temp" in fn
    assert "for update" in fn.split("category = 'send_out'")[1][:40]
    assert "from auto_send_runs where id = v_run_id for update" in fn
    for code in (
        "project_not_found",
        "run_not_for_project",
        "run_not_sent",
        "already_accepted",
        "already_sent",
        "external_submission_active",
        "outcome_recorded",
        "stale",
        "gc_already_sent",
        "no_sent_gcs",
    ):
        assert f"'auto_send_accept:{code}'" in fn, code
    assert "'auto_send'" in fn and "auto_send_run_id" in fn
    assert "where proposal_sends.status not in ('sent', 'sending')" in fn
    assert "insert into proposal_send_events" in fn
    # cc_recipients is same-GC contact objects (0110); the bids desk CC that
    # auto_send_gcs.cc holds is never copied there.
    assert "v_g.cc" not in fn
    assert "insert into stage_events" in fn and "coalesce(v_run.sent_at, now())" in fn
    assert "current_stage = 'submitted'" in fn
    assert "set accepted_at = now(), accepted_by = v_actor" in fn
    for grant in (
        "revoke all on function apply_auto_send_acceptance(jsonb) from public;",
        "revoke all on function apply_auto_send_acceptance(jsonb) from anon, authenticated;",
        "grant execute on function apply_auto_send_acceptance(jsonb) to service_role;",
    ):
        assert grant in sql
