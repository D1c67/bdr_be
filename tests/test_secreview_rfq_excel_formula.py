"""Security review: the RFQ workbook letterhead and category banner must not
render a user-typed leading '=' as a live formula (CWE-1236)."""

import io

import openpyxl

from app.services.rfq_excel import build_rfq_workbook


def _sheet(category: str, project: dict | None):
    data = build_rfq_workbook(category, [{"sr_no": 1, "description": "Wire", "quantity": 1}], project)
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_category_banner_formula_is_neutralized():
    ws = _sheet('=HYPERLINK("evil.example","x")', {"number": "26.9.0001", "name": "Job"})
    cell = ws["A7"]
    assert cell.data_type == "s"
    assert str(cell.value).startswith("'=")


def test_project_label_without_number_is_neutralized():
    ws = _sheet("Lighting", {"number": None, "name": "=cmd|/C calc!A0"})
    cell = ws["B2"]
    assert cell.data_type == "s"
    assert cell.value == "'=cmd|/C calc!A0"


def test_plain_headers_unchanged():
    ws = _sheet("Lighting", {"number": "26.9.0001", "name": "Job"})
    assert ws["A7"].value == "LIGHTING"
    assert ws["B2"].value == "26.9.0001 - Job"
    assert ws["B1"].value == "REQUEST FOR QUOTE"
